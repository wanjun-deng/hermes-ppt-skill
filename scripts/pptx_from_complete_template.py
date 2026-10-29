#!/usr/bin/env python3
"""Execute a ``powerpoint.plan/v1`` plan against a complete PPTX template.

The plan is declarative and sequential: the caller selects source slides,
layouts, targets, and content while this script performs deterministic PPTX
mutations. Original slides use the ``recommended_ref`` emitted by
``pptx_read.py inspect``; newly duplicated or appended slides are addressed
through caller-defined aliases. Layout and shape locators are likewise strict
inventory references. Table dimensions must match the template exactly, and
slides containing charts cannot be duplicated.
"""
import argparse
import json
import os
import re
import sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Pt
from PIL import Image

from pptx_common import PptxSkillError, atomic_save_presentation
from pptx_create import add_bullets
from pptx_edit import (R_EMBED, duplicate_slide, enable_slide_number,
                       move_slide, remove_slide, set_background, set_footer,
                       set_notes)
from pptx_inventory import (InventoryError, is_picture_shape, node_for_shape, recommended_ref,
                            resolve_shape_ref, sha256_file, walk_shapes)
from pptx_text import replace_all_text, replace_text


class PlanError(PptxSkillError):
    """A user-correctable generation plan error."""

    def __init__(self, message, code="invalid_plan", operation_index=None):
        super().__init__(message)
        self.code = code
        self.operation_index = operation_index

    def as_error(self):
        result = {"code": self.code, "message": str(self)}
        if self.operation_index is not None:
            result["operation_index"] = self.operation_index
        return result


def _shapes(slide):
    for node in walk_shapes(slide.shapes):
        yield node.shape


def _shape_ref(operation):
    ref = operation.get("shape")
    if ref is None:
        ref = operation.get("target")
    return ref


def _operation_ref(operation, payload, slide, allow_inferred_roles):
    """Resolve an explicit shape ref or a semantic role binding."""
    ref = _shape_ref(operation)
    if ref is None and operation.get("role") is not None:
        targets = payload.get("targets", {})
        if not isinstance(targets, dict):
            raise PlanError("slide targets must be an object")
        ref = targets.get(operation["role"])
        if ref is None and not allow_inferred_roles:
            raise PlanError(
                f"role {operation['role']!r} requires an explicit target binding",
                code="unbound_role")
        if ref is None:
            ref = infer_role_target(slide, operation["role"])
        if ref is None:
            raise PlanError(
                f"no target binding for role {operation['role']!r}")
    return ref


def _text_shapes(slide):
    return [shape for shape in _shapes(slide)
            if shape.has_text_frame and shape.text.strip()]


def _font_score(shape):
    sizes = [run.font.size.pt for para in shape.text_frame.paragraphs
             for run in para.runs if run.font.size is not None]
    return max(sizes, default=0)


def _shape_locator(slide, shape):
    return recommended_ref(node_for_shape(slide, shape))


def infer_role_target(slide, role):
    """Infer stable common roles from the template's existing text."""
    shapes = _text_shapes(slide)
    if not shapes:
        return None
    if role == "date":
        match = next((s for s in shapes
                      if re.search(r"(汇报时间|日期|时间|date)", s.text,
                                   re.IGNORECASE)), None)
        return _shape_locator(slide, match) if match is not None else None
    if role == "presenter":
        match = next((s for s in shapes
                      if re.search(r"(汇报人|作者|负责人|presenter|author)",
                                   s.text, re.IGNORECASE)), None)
        return _shape_locator(slide, match) if match is not None else None
    if role == "period":
        match = next((s for s in shapes
                      if re.search(r"(年中|年终|季度|月度|q[1-4])", s.text,
                                   re.IGNORECASE)), None)
        return _shape_locator(slide, match) if match is not None else None
    if role == "subtitle":
        candidates = [s for s in shapes if re.fullmatch(
            r"[A-Z0-9][A-Z0-9 .&:/-]{3,}", s.text.strip())]
        match = max(candidates, key=_font_score) if candidates else None
        return _shape_locator(slide, match) if match is not None else None
    if role == "title":
        match = max(
            shapes,
            key=lambda shape: (
                _font_score(shape),
                -node_for_shape(slide, shape).slide_bounds[1],
            ),
        )
        return _shape_locator(slide, match)
    if role == "summary":
        match = max(shapes, key=lambda s: len(s.text))
        return _shape_locator(slide, match)
    return None


def find_shape(slide, ref):
    """Resolve a strict locator emitted by pptx_read.py inspect."""
    try:
        return resolve_shape_ref(slide, ref)
    except InventoryError as exc:
        raise PlanError(str(exc)) from exc


def _operations(value):
    if value is None:
        return []
    return value if isinstance(value, list) else [value]


def fill_text(shape, value):
    """Replace a shape's text using the shared format-preserving policy."""
    if not shape.has_text_frame:
        raise PlanError(f"shape {shape.name!r} has no text frame")
    replace_all_text(shape.text_frame, value)


def fill_table(shape, rows):
    if not shape.has_table:
        raise PlanError(f"shape {shape.name!r} is not a table")
    table = shape.table
    if len(rows) != len(table.rows) or any(
            len(row) != len(table.columns) for row in rows):
        raise PlanError(
            f"table {shape.name!r} dimensions do not match template: "
            f"expected {len(table.rows)}x{len(table.columns)}, "
            f"got {len(rows)}x{max((len(row) for row in rows), default=0)}")
    for row, cells in zip(table.rows, rows):
        for cell, value in zip(row.cells, cells):
            replace_all_text(cell.text_frame, value)


def fill_chart(shape, spec):
    if not shape.has_chart:
        raise PlanError(f"shape {shape.name!r} is not a chart")
    categories = spec.get("categories")
    series = spec.get("series")
    if not isinstance(categories, list) or not categories:
        raise PlanError(f"chart {shape.name!r} requires non-empty categories")
    if not isinstance(series, dict) or not series:
        raise PlanError(f"chart {shape.name!r} requires non-empty series")
    category_count = len(categories)
    for name, values in series.items():
        if not isinstance(values, list) or len(values) != category_count:
            raise PlanError(
                f"chart {shape.name!r} series {name!r} must contain "
                f"{category_count} values")
    data = CategoryChartData()
    data.categories = categories
    for name, values in series.items():
        data.add_series(name, values)
    shape.chart.replace_data(data)
    if "number_format" in spec:
        axis = shape.chart.value_axis
        axis.tick_labels.number_format = spec["number_format"]
        axis.tick_labels.number_format_is_linked = False
        for plot in shape.chart.plots:
            if plot.has_data_labels:
                plot.data_labels.number_format = spec["number_format"]
                plot.data_labels.number_format_is_linked = False
    if spec.get("display_units") == "none":
        axis = shape.chart.value_axis
        for element in list(axis._element.findall(qn("c:dispUnits"))):
            axis._element.remove(element)
    if spec.get("title") is not None:
        shape.chart.has_title = True
        replace_all_text(shape.chart.chart_title.text_frame, spec["title"])


def replace_picture(shape, path, fit="preserve"):
    empty_placeholder = (shape.is_placeholder and not is_picture_shape(shape)
                         and shape.placeholder_format.type == PP_PLACEHOLDER.PICTURE)
    if not is_picture_shape(shape) and not empty_placeholder:
        raise PlanError(f"shape {shape.name!r} is not a picture")
    if not os.path.isfile(path):
        raise PlanError(f"image file not found: {path}")
    if empty_placeholder:
        shape = shape.insert_picture(path)
    else:
        _, rid = shape.part.get_or_add_image_part(path)
        shape._element.blipFill.blip.set(R_EMBED + "embed", rid)
    if fit != "preserve":
        left, top, width, height = shape.left, shape.top, shape.width, shape.height
        with Image.open(path) as image:
            ratio = image.width / image.height
        shape.crop_left = shape.crop_right = shape.crop_top = shape.crop_bottom = 0
        box_ratio = width / height
        # A newly filled placeholder may inherit all geometry from its layout.
        # Materialize the entire transform before changing one dimension;
        # otherwise python-pptx creates zero extents for the inherited axes.
        shape.left, shape.top, shape.width, shape.height = left, top, width, height
        if fit == "cover":
            if ratio > box_ratio:
                shape.crop_left = shape.crop_right = (1 - box_ratio / ratio) / 2
            else:
                shape.crop_top = shape.crop_bottom = (1 - ratio / box_ratio) / 2
        elif fit == "contain":
            if ratio > box_ratio:
                shape.height = round(width / ratio)
                shape.top = top + (height - shape.height) // 2
            else:
                shape.width = round(height * ratio)
                shape.left = left + (width - shape.width) // 2


def apply_payload(slide, payload, report, allow_inferred_roles=False):
    for operation_number, operation in enumerate(
            _operations(payload.get("text")), start=1):
        ref = _operation_ref(
            operation, payload, slide, allow_inferred_roles)
        if ref is None:
            raise PlanError(
                f"text operation {operation_number} requires shape or target")
        try:
            shape = find_shape(slide, ref)
        except PlanError as exc:
            raise PlanError(
                f"text operation {operation_number} failed: {exc}") from exc
        if operation.get("required", True) and not str(
                operation.get("value", "")).strip():
            raise PlanError(
                f"required text operation {operation_number} has an empty value")
        fill_text(shape, operation.get("value", ""))
        if "font_size_pt" in operation:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(operation["font_size_pt"])
        if "alignment" in operation:
            for paragraph in shape.text_frame.paragraphs:
                paragraph.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[operation["alignment"]]
        if operation.get("auto_size") == "none":
            shape.text_frame.auto_size = MSO_AUTO_SIZE.NONE
        if "word_wrap" in operation:
            shape.text_frame.word_wrap = operation["word_wrap"]
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if "font_family" in operation:
                    from pptx_fonts import set_run_fonts
                    set_run_fonts(run, latin=operation["font_family"],
                                  east_asian=operation["font_family"],
                                  complex_script=operation["font_family"])
                if "character_spacing_pt" in operation:
                    run._r.get_or_add_rPr().set("spc", str(round(operation["character_spacing_pt"] * 100)))
        report["filled_text"] += 1

    for operation in _operations(payload.get("images")):
        replace_picture(find_shape(slide, _shape_ref(operation)),
                        operation["path"], operation.get("fit", "preserve"))
        report["replaced_images"] += 1

    for operation in _operations(payload.get("tables")):
        fill_table(find_shape(slide, _shape_ref(operation)), operation["rows"])
        report["updated_tables"] += 1

    for operation in _operations(payload.get("charts")):
        fill_chart(find_shape(slide, _shape_ref(operation)), operation)
        report["updated_charts"] += 1

    if payload.get("title") is not None:
        if slide.shapes.title is None:
            raise PlanError("slide has no title placeholder")
        replace_all_text(slide.shapes.title.text_frame, payload["title"])
        report["filled_text"] += 1
    if payload.get("bullets"):
        body = next((
            ph for ph in slide.placeholders
            if ph.placeholder_format.type in {
                PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT,
            }
        ), None)
        if body is None:
            raise PlanError("slide has no body placeholder for bullets")
        add_bullets(body.text_frame, payload["bullets"])
        report["filled_text"] += 1
    if payload.get("notes") is not None:
        set_notes(slide, str(payload["notes"]))
        report["updated_notes"] += 1
    if payload.get("background"):
        set_background(slide, payload["background"])
        report["updated_backgrounds"] += 1
    if payload.get("footer") is not None:
        set_footer(slide, str(payload["footer"]))
        report["updated_footers"] += 1
    if payload.get("slide_number"):
        if not enable_slide_number(slide):
            raise PlanError("layout provides no slide-number placeholder")
        report["enabled_slide_numbers"] += 1


def find_layout(prs, ref):
    if isinstance(ref, int):
        try:
            return prs.slide_layouts[ref]
        except IndexError as exc:
            raise PlanError(f"layout index not found: {ref}") from exc
    for layout in prs.slide_layouts:
        if layout.name == ref:
            return layout
    raise PlanError(f"layout not found: {ref!r}")


PLAN_SCHEMA = "powerpoint.plan/v1"
RESULT_SCHEMA = "powerpoint.plan-result/v1"
SHAPE_REF_FIELDS = {"placeholder", "shape_id", "path", "name"}
PAYLOAD_FIELDS = {
    "text", "images", "tables", "charts", "targets", "title", "bullets",
    "notes", "background", "footer", "slide_number",
}


def _reject_unknown(mapping, allowed, context):
    unknown = sorted(set(mapping) - set(allowed))
    if unknown:
        raise PlanError(
            f"{context} contains unknown fields: {unknown}",
            code="unknown_field")


def _require_int(value, context, minimum=0):
    if isinstance(value, bool) or not isinstance(value, int) or value < minimum:
        raise PlanError(
            f"{context} must be an integer >= {minimum}",
            code="invalid_type")
    return value


def _validate_shape_ref(ref, context):
    if not isinstance(ref, dict):
        raise PlanError(f"{context} must be an object", code="invalid_shape_ref")
    _reject_unknown(ref, SHAPE_REF_FIELDS, context)
    if not ref:
        raise PlanError(f"{context} cannot be empty", code="invalid_shape_ref")
    if not (set(ref) & SHAPE_REF_FIELDS):
        raise PlanError(f"{context} has no locator field", code="invalid_shape_ref")
    if "placeholder" in ref:
        _require_int(ref["placeholder"], f"{context}.placeholder")
    if "shape_id" in ref:
        _require_int(ref["shape_id"], f"{context}.shape_id", minimum=1)
    if "path" in ref:
        path = ref["path"]
        if not isinstance(path, list) or not path:
            raise PlanError(
                f"{context}.path must be a non-empty array",
                code="invalid_shape_ref")
        for position, value in enumerate(path):
            _require_int(value, f"{context}.path[{position}]")
    if "name" in ref and not isinstance(ref["name"], str):
        raise PlanError(f"{context}.name must be a string", code="invalid_type")


def _validate_slide_ref(ref, context, known_aliases):
    if not isinstance(ref, dict):
        raise PlanError(f"{context} must be an object", code="invalid_slide_ref")
    if set(ref) == {"alias"}:
        alias = ref["alias"]
        if not isinstance(alias, str) or not alias:
            raise PlanError(
                f"{context}.alias must be a non-empty string",
                code="invalid_slide_ref")
        if alias not in known_aliases:
            raise PlanError(
                f"{context} references alias {alias!r} before it is declared",
                code="unknown_slide_alias")
        return
    if set(ref) != {"slide_id", "index"}:
        raise PlanError(
            f"{context} must contain exactly slide_id+index or alias",
            code="invalid_slide_ref")
    _require_int(ref["slide_id"], f"{context}.slide_id", minimum=1)
    _require_int(ref["index"], f"{context}.index")


def _validate_layout_ref(ref, context):
    if not isinstance(ref, dict) or set(ref) != {"index", "name"}:
        raise PlanError(
            f"{context} must contain exactly index and name",
            code="invalid_layout_ref")
    _require_int(ref["index"], f"{context}.index")
    if not isinstance(ref["name"], str) or not ref["name"]:
        raise PlanError(
            f"{context}.name must be a non-empty string",
            code="invalid_layout_ref")


def _validate_operation_list(value, field, operation_number):
    context = f"operations[{operation_number}].{field}"
    if not isinstance(value, list):
        raise PlanError(f"{context} must be an array", code="invalid_type")
    for item_number, item in enumerate(value):
        if not isinstance(item, dict):
            raise PlanError(
                f"{context}[{item_number}] must be an object",
                code="invalid_type")
    return value


def _validate_payload(operation, operation_number):
    prefix = f"operations[{operation_number}]"
    targets = operation.get("targets")
    if targets is not None:
        if not isinstance(targets, dict):
            raise PlanError(f"{prefix}.targets must be an object", code="invalid_type")
        for role, ref in targets.items():
            if not isinstance(role, str) or not role:
                raise PlanError(
                    f"{prefix}.targets keys must be non-empty strings",
                    code="invalid_type")
            _validate_shape_ref(ref, f"{prefix}.targets[{role!r}]")

    if "text" in operation:
        for item_number, item in enumerate(_validate_operation_list(
                operation["text"], "text", operation_number)):
            context = f"{prefix}.text[{item_number}]"
            _reject_unknown(item, {"shape", "role", "value", "required", "font_size_pt", "alignment", "auto_size", "word_wrap", "font_family", "character_spacing_pt"}, context)
            if "font_family" in item and (not isinstance(item["font_family"], str) or not item["font_family"].strip()):
                raise PlanError(f"{context}.font_family must be a nonempty string", code="invalid_type")
            if "character_spacing_pt" in item and (type(item["character_spacing_pt"]) not in (int, float) or not -5 <= item["character_spacing_pt"] <= 20):
                raise PlanError(f"{context}.character_spacing_pt must be -5..20", code="invalid_type")
            if "auto_size" in item and item["auto_size"] != "none":
                raise PlanError(f"{context}.auto_size supports only none", code="invalid_type")
            if "word_wrap" in item and not isinstance(item["word_wrap"], bool):
                raise PlanError(f"{context}.word_wrap must be boolean", code="invalid_type")
            if "alignment" in item and item["alignment"] not in {"left", "center", "right"}:
                raise PlanError(f"{context}.alignment must be left, center or right", code="invalid_type")
            if "font_size_pt" in item and (type(item["font_size_pt"]) not in (int, float) or not 1 <= item["font_size_pt"] <= 400):
                raise PlanError(f"{context}.font_size_pt must be 1..400", code="invalid_type")
            if ("shape" in item) == ("role" in item):
                raise PlanError(
                    f"{context} requires exactly one of shape or role",
                    code="invalid_text_operation")
            if "shape" in item:
                _validate_shape_ref(item["shape"], f"{context}.shape")
            if "role" in item and (not isinstance(item["role"], str)
                                    or not item["role"]):
                raise PlanError(
                    f"{context}.role must be a non-empty string",
                    code="invalid_type")
            if "value" not in item:
                raise PlanError(f"{context} requires value", code="missing_field")
            if "required" in item and not isinstance(item["required"], bool):
                raise PlanError(
                    f"{context}.required must be boolean", code="invalid_type")

    if "images" in operation:
        for item_number, item in enumerate(_validate_operation_list(
                operation["images"], "images", operation_number)):
            context = f"{prefix}.images[{item_number}]"
            _reject_unknown(item, {"shape", "path", "fit"}, context)
            if not {"shape", "path"} <= set(item):
                raise PlanError(
                    f"{context} requires shape and path", code="missing_field")
            _validate_shape_ref(item["shape"], f"{context}.shape")
            if item.get("fit", "preserve") not in {"preserve", "cover", "contain"}:
                raise PlanError(f"{context}.fit must be preserve, cover or contain", code="invalid_type")
            if not isinstance(item["path"], str) or not item["path"]:
                raise PlanError(
                    f"{context}.path must be a non-empty string",
                    code="invalid_type")

    if "tables" in operation:
        for item_number, item in enumerate(_validate_operation_list(
                operation["tables"], "tables", operation_number)):
            context = f"{prefix}.tables[{item_number}]"
            _reject_unknown(item, {"shape", "rows"}, context)
            if set(item) != {"shape", "rows"}:
                raise PlanError(
                    f"{context} requires shape and rows", code="missing_field")
            _validate_shape_ref(item["shape"], f"{context}.shape")
            rows = item["rows"]
            if (not isinstance(rows, list) or not rows
                    or any(not isinstance(row, list) for row in rows)):
                raise PlanError(
                    f"{context}.rows must be a non-empty array of arrays",
                    code="invalid_table")

    if "charts" in operation:
        for item_number, item in enumerate(_validate_operation_list(
                operation["charts"], "charts", operation_number)):
            context = f"{prefix}.charts[{item_number}]"
            _reject_unknown(item, {"shape", "categories", "series", "title", "number_format", "display_units"}, context)
            if "number_format" in item and (not isinstance(item["number_format"], str) or not item["number_format"]):
                raise PlanError(f"{context}.number_format must be a nonempty string", code="invalid_type")
            if "display_units" in item and item["display_units"] != "none":
                raise PlanError(f"{context}.display_units supports only none", code="invalid_type")
            required = {"shape", "categories", "series"}
            if not required <= set(item):
                raise PlanError(
                    f"{context} requires shape, categories, and series",
                    code="missing_field")
            _validate_shape_ref(item["shape"], f"{context}.shape")
            categories = item["categories"]
            series = item["series"]
            if not isinstance(categories, list) or not categories:
                raise PlanError(
                    f"{context}.categories must be a non-empty array",
                    code="invalid_chart")
            if not isinstance(series, dict) or not series:
                raise PlanError(
                    f"{context}.series must be a non-empty object",
                    code="invalid_chart")
            for name, values in series.items():
                if (not isinstance(name, str) or not isinstance(values, list)
                        or len(values) != len(categories)):
                    raise PlanError(
                        f"{context}.series entries must match category count",
                        code="invalid_chart")

    if "bullets" in operation:
        bullets = operation["bullets"]
        if not isinstance(bullets, list):
            raise PlanError(
                f"{prefix}.bullets must be an array", code="invalid_type")
        for item_number, item in enumerate(bullets):
            context = f"{prefix}.bullets[{item_number}]"
            if isinstance(item, str):
                continue
            if not isinstance(item, dict):
                raise PlanError(
                    f"{context} must be a string or object",
                    code="invalid_type")
            _reject_unknown(
                item,
                {"text", "level", "size", "bold", "italic", "font", "color",
                 "link"},
                context,
            )
            if "text" not in item or not isinstance(item["text"], str):
                raise PlanError(
                    f"{context}.text must be a string", code="missing_field")
            if "level" in item:
                level = item["level"]
                if (isinstance(level, bool) or not isinstance(level, int)
                        or level not in range(5)):
                    raise PlanError(
                        f"{context}.level must be 0..4",
                        code="invalid_value")
            for field in ("bold", "italic"):
                if field in item and not isinstance(item[field], bool):
                    raise PlanError(
                        f"{context}.{field} must be boolean",
                        code="invalid_type")
            for field in ("size",):
                if field in item and (
                        isinstance(item[field], bool)
                        or not isinstance(item[field], (int, float))
                        or item[field] <= 0):
                    raise PlanError(
                        f"{context}.{field} must be a positive number",
                        code="invalid_value")
            for field in ("font", "color", "link"):
                if field in item and not isinstance(item[field], str):
                    raise PlanError(
                        f"{context}.{field} must be a string",
                        code="invalid_type")
            if ("color" in item and not re.fullmatch(
                    r"[0-9A-Fa-f]{6}", item["color"])):
                raise PlanError(
                    f"{context}.color must be a six-digit RGB color",
                    code="invalid_color")
    for field in ("title", "notes", "footer"):
        if field in operation and not isinstance(operation[field], str):
            raise PlanError(
                f"{prefix}.{field} must be a string", code="invalid_type")
    if "background" in operation:
        color = operation["background"]
        if not isinstance(color, str) or not re.fullmatch(r"[0-9A-Fa-f]{6}", color):
            raise PlanError(
                f"{prefix}.background must be a six-digit RGB color",
                code="invalid_color")
    if "slide_number" in operation and not isinstance(
            operation["slide_number"], bool):
        raise PlanError(
            f"{prefix}.slide_number must be boolean", code="invalid_type")


def _bind_initial_slide(prs, ref, context):
    index = ref["index"]
    if index >= len(prs.slides):
        raise PlanError(f"{context}.index is out of range: {index}",
                        code="invalid_slide_ref")
    actual = int(prs.slides[index].slide_id)
    if actual != ref["slide_id"]:
        raise PlanError(
            f"{context} slide_id mismatch: expected {ref['slide_id']}, found {actual}",
            code="invalid_slide_ref")
    return actual


def _bind_layout(prs, ref, context):
    index = ref["index"]
    if index >= len(prs.slide_layouts):
        raise PlanError(f"{context}.index is out of range: {index}",
                        code="invalid_layout_ref")
    actual = prs.slide_layouts[index].name
    if actual != ref["name"]:
        raise PlanError(
            f"{context} layout name mismatch: expected {ref['name']!r}, found {actual!r}",
            code="invalid_layout_ref")
    return index


def compile_plan(prs, plan, template_path):
    """Validate the v1 plan and bind original slide/layout references."""
    if not isinstance(plan, dict):
        raise PlanError("plan root must be an object", code="invalid_type")
    _reject_unknown(
        plan,
        {"schema", "template", "strict", "allow_inferred_roles", "values",
         "operations"},
        "plan")
    if plan.get("schema") != PLAN_SCHEMA:
        raise PlanError(
            f"plan.schema must be {PLAN_SCHEMA!r}", code="unsupported_schema")
    template_meta = plan.get("template")
    if not isinstance(template_meta, dict):
        raise PlanError("plan.template must be an object", code="missing_field")
    _reject_unknown(template_meta, {"sha256", "slide_count"}, "plan.template")
    if set(template_meta) != {"sha256", "slide_count"}:
        raise PlanError(
            "plan.template requires sha256 and slide_count", code="missing_field")
    expected_hash = template_meta["sha256"]
    if not isinstance(expected_hash, str) or not re.fullmatch(
            r"[0-9A-Fa-f]{64}", expected_hash):
        raise PlanError("plan.template.sha256 must be 64 hexadecimal characters",
                        code="invalid_template_identity")
    if expected_hash.lower() != sha256_file(template_path).lower():
        raise PlanError("template sha256 mismatch", code="template_mismatch")
    expected_count = _require_int(
        template_meta["slide_count"], "plan.template.slide_count")
    if expected_count != len(prs.slides):
        raise PlanError(
            f"template slide count mismatch: expected {expected_count}, got {len(prs.slides)}",
            code="template_mismatch")
    if "strict" in plan and not isinstance(plan["strict"], bool):
        raise PlanError("plan.strict must be boolean", code="invalid_type")
    if "allow_inferred_roles" in plan and not isinstance(
            plan["allow_inferred_roles"], bool):
        raise PlanError("plan.allow_inferred_roles must be boolean",
                        code="invalid_type")
    values = plan.get("values", {})
    if not isinstance(values, dict) or any(not isinstance(key, str) for key in values):
        raise PlanError("plan.values must be an object with string keys",
                        code="invalid_type")
    operations = plan.get("operations")
    if not isinstance(operations, list):
        raise PlanError("plan.operations must be an array", code="missing_field")

    compiled = []
    known_aliases = set()
    for number, operation in enumerate(operations):
        try:
            if not isinstance(operation, dict):
                raise PlanError(
                    f"operations[{number}] must be an object", code="invalid_type",
                    operation_index=number)
            action = operation.get("action")
            allowed_by_action = {
                "fill": {"action", "slide"} | PAYLOAD_FIELDS,
                "duplicate": {"action", "slide", "as"} | PAYLOAD_FIELDS,
                "append": {"action", "layout", "as"} | PAYLOAD_FIELDS,
                "remove": {"action", "slide"},
                "reorder": {"action", "slide", "to"},
            }
            if action not in allowed_by_action:
                raise PlanError(
                    f"operations[{number}].action is invalid: {action!r}",
                    code="unknown_action", operation_index=number)
            _reject_unknown(operation, allowed_by_action[action],
                            f"operations[{number}]")
            compiled_operation = dict(operation)

            if action in {"fill", "duplicate", "remove", "reorder"}:
                if "slide" not in operation:
                    raise PlanError(
                        f"operations[{number}] requires slide", code="missing_field",
                        operation_index=number)
                ref = operation["slide"]
                _validate_slide_ref(ref, f"operations[{number}].slide", known_aliases)
                if "alias" not in ref:
                    compiled_operation["_bound_slide_id"] = _bind_initial_slide(
                        prs, ref, f"operations[{number}].slide")

            if action == "append":
                if "layout" not in operation:
                    raise PlanError(
                        f"operations[{number}] requires layout", code="missing_field",
                        operation_index=number)
                _validate_layout_ref(operation["layout"],
                                     f"operations[{number}].layout")
                compiled_operation["_bound_layout_index"] = _bind_layout(
                    prs, operation["layout"], f"operations[{number}].layout")

            if action in {"fill", "duplicate", "append"}:
                _validate_payload(operation, number)

            if action == "reorder":
                if "to" not in operation:
                    raise PlanError(
                        f"operations[{number}] requires to", code="missing_field",
                        operation_index=number)
                _require_int(operation["to"], f"operations[{number}].to")

            alias = operation.get("as")
            if alias is not None:
                if not isinstance(alias, str) or not alias:
                    raise PlanError(
                        f"operations[{number}].as must be a non-empty string",
                        code="invalid_slide_alias", operation_index=number)
                if alias in known_aliases:
                    raise PlanError(
                        f"duplicate slide alias: {alias!r}", code="duplicate_slide_alias",
                        operation_index=number)
                known_aliases.add(alias)
            compiled.append(compiled_operation)
        except PlanError as exc:
            if exc.operation_index is None:
                exc.operation_index = number
            raise
    return compiled


def _slide_by_id(prs, slide_id, operation_index):
    for index, slide in enumerate(prs.slides):
        if int(slide.slide_id) == int(slide_id):
            return index, slide
    raise PlanError(
        f"slide_id {slide_id} is no longer present",
        code="invalid_slide_ref", operation_index=operation_index)


def _runtime_slide(prs, operation, aliases, operation_index):
    ref = operation["slide"]
    if "alias" in ref:
        alias = ref["alias"]
        if alias not in aliases:
            raise PlanError(
                f"slide alias {alias!r} is not available",
                code="unknown_slide_alias", operation_index=operation_index)
        slide_id = aliases[alias]
    else:
        slide_id = operation["_bound_slide_id"]
    return _slide_by_id(prs, slide_id, operation_index)


def execute_plan(template, plan, output):
    prs = Presentation(template)
    operations = compile_plan(prs, plan, template)
    counters = {
        "filled_text": 0, "replaced_images": 0, "updated_tables": 0,
        "updated_charts": 0, "updated_notes": 0, "updated_backgrounds": 0,
        "updated_footers": 0, "enabled_slide_numbers": 0,
        "duplicated_slides": 0, "removed_slides": 0,
        "appended_slides": 0, "reordered_slides": 0,
    }
    report = {
        "schema": RESULT_SCHEMA,
        "ok": True,
        "output": output,
        "template_sha256": sha256_file(template),
        "slides_before": len(prs.slides),
        "slides_after": None,
        "operations": counters,
        "applied": [],
        "aliases": {},
        "warnings": [],
    }
    for token, value in plan.get("values", {}).items():
        replacements = replace_text(prs, "{{%s}}" % token, str(value))
        counters["filled_text"] += replacements
        if plan.get("strict", False) and replacements == 0:
            raise PlanError(
                f"value token not found in template: {{{{{token}}}}}",
                code="token_not_found")

    aliases = {}
    allow_inferred_roles = plan.get("allow_inferred_roles", False)
    for number, operation in enumerate(operations):
        action = operation["action"]
        applied = {"index": number, "action": action}
        try:
            if action == "fill":
                _, slide = _runtime_slide(prs, operation, aliases, number)
                apply_payload(slide, operation, counters,
                              allow_inferred_roles)
                applied["slide_id"] = int(slide.slide_id)
            elif action == "duplicate":
                index, source = _runtime_slide(prs, operation, aliases, number)
                if any(shape.has_chart for shape in _shapes(source)):
                    raise PlanError(
                        f"cannot duplicate chart slide {index}",
                        code="unsupported_chart_duplication")
                new_index = duplicate_slide(prs, index)
                slide = prs.slides[new_index]
                apply_payload(slide, operation, counters,
                              allow_inferred_roles)
                counters["duplicated_slides"] += 1
                applied["slide_id"] = int(slide.slide_id)
                applied["created_index"] = new_index
                if operation.get("as"):
                    aliases[operation["as"]] = int(slide.slide_id)
            elif action == "append":
                layout = prs.slide_layouts[operation["_bound_layout_index"]]
                slide = prs.slides.add_slide(layout)
                apply_payload(slide, operation, counters,
                              allow_inferred_roles)
                counters["appended_slides"] += 1
                applied["slide_id"] = int(slide.slide_id)
                applied["created_index"] = len(prs.slides) - 1
                if operation.get("as"):
                    aliases[operation["as"]] = int(slide.slide_id)
            elif action == "remove":
                index, slide = _runtime_slide(prs, operation, aliases, number)
                applied["slide_id"] = int(slide.slide_id)
                remove_slide(prs, index)
                counters["removed_slides"] += 1
            elif action == "reorder":
                index, slide = _runtime_slide(prs, operation, aliases, number)
                destination = operation["to"]
                if destination >= len(prs.slides):
                    raise PlanError(
                        f"reorder destination is out of range: {destination}",
                        code="invalid_slide_ref")
                move_slide(prs, index, destination)
                counters["reordered_slides"] += 1
                applied["slide_id"] = int(slide.slide_id)
                applied["to"] = destination
        except PlanError as exc:
            if exc.operation_index is None:
                exc.operation_index = number
            raise
        except PptxSkillError as exc:
            raise PlanError(str(exc), code=exc.code,
                            operation_index=number) from exc
        report["applied"].append(applied)

    output = atomic_save_presentation(prs, output)
    report["output"] = output
    report["slides_after"] = len(prs.slides)
    for alias, slide_id in aliases.items():
        try:
            index, _ = _slide_by_id(prs, slide_id, None)
            report["aliases"][alias] = {"slide_id": slide_id, "index": index}
        except PlanError:
            report["aliases"][alias] = {"slide_id": slide_id, "removed": True}
    return report

def _write_error(error):
    print(json.dumps({
        "schema": "powerpoint.error/v1",
        "ok": False,
        "error": error,
    }, ensure_ascii=False), file=sys.stderr)


def main(argv=None):
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser(
        description="Generate a PowerPoint from a complex template and plan.")
    parser.add_argument("template", help="template .pptx path")
    parser.add_argument("plan", help="powerpoint.plan/v1 JSON path")
    parser.add_argument("output", help="output .pptx path")
    args = parser.parse_args(argv)
    try:
        with open(args.plan, encoding="utf-8") as fh:
            plan = json.load(fh)
        result = execute_plan(args.template, plan, args.output)
    except PlanError as exc:
        _write_error(exc.as_error())
        return 2
    except InventoryError as exc:
        _write_error({"code": "unsafe_inventory", "message": str(exc)})
        return 2
    except PptxSkillError as exc:
        _write_error(exc.as_error())
        return 2
    except (OSError, KeyError, TypeError, ValueError, AttributeError,
            IndexError) as exc:
        _write_error({"code": "execution_failed", "message": str(exc)})
        return 2
    print(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    sys.exit(main())
