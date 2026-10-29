#!/usr/bin/env python3
"""Shared formatting-preserving text traversal and mutation helpers."""

import copy

from lxml import etree
from pptx.oxml.ns import qn

from pptx_inventory import walk_shapes


def _run_format_key(run_element):
    properties = run_element.find(qn("a:rPr"))
    return None if properties is None else etree.tostring(properties)


def normalize_runs(paragraph):
    """Merge adjacent runs whose formatting XML is byte-identical."""
    runs = list(paragraph.runs)
    index = 0
    while index + 1 < len(runs):
        first, second = runs[index], runs[index + 1]
        if (_run_format_key(first._r) == _run_format_key(second._r)
                and first._r.getnext() is second._r):
            first.text += second.text
            second._r.getparent().remove(second._r)
            runs.pop(index + 1)
        else:
            index += 1


def replace_in_text_frame(text_frame, old, new):
    """Keep run styles/links; inserted text takes the first matched style.

    Fields and explicit breaks are boundaries. Unmatched characters retain their
    original runs, including suffixes after matches spanning mixed formatting.
    """
    if old == "":
        raise ValueError("replacement source text cannot be empty")
    count = 0
    for paragraph in text_frame.paragraphs:
        by_element = {run._r: run for run in paragraph.runs}
        sequences, current = [], []
        for element in paragraph._p:
            if element.tag == qn("a:r"):
                current.append(by_element[element])
            elif element.tag in {qn("a:br"), qn("a:fld")}:
                sequences.append(current)
                current = []
        sequences.append(current)
        for runs in sequences:
            joined = "".join(run.text for run in runs)
            starts, ends, offset = [], [], 0
            for run in runs:
                starts.append(offset)
                offset += len(run.text)
                ends.append(offset)
            matches, position = [], 0
            while (position := joined.find(old, position)) >= 0:
                matches.append(position)
                position += len(old)
            for start in reversed(matches):
                end = start + len(old)
                first = next(i for i in range(len(runs)) if starts[i] <= start < ends[i])
                last = next(i for i in range(first, len(runs)) if starts[i] < end <= ends[i])
                prefix = runs[first].text[:start - starts[first]]
                suffix = runs[last].text[end - starts[last]:]
                if first == last:
                    runs[first].text = prefix + new + suffix
                else:
                    runs[first].text = prefix + new
                    for i in range(first + 1, last):
                        runs[i].text = ""
                    runs[last].text = suffix
            count += len(matches)
    return count


def replace_all_text(text_frame, value):
    """Replace a whole text frame while preserving its first styled paragraph."""
    text = str(value)
    body = text_frame._txBody
    template = next(
        (paragraph for paragraph in body.findall(qn("a:p"))
         if paragraph.findall(qn("a:r"))),
        None,
    )
    if template is None:
        text_frame.text = text
        return
    for paragraph in body.findall(qn("a:p")):
        body.remove(paragraph)
    for line in text.split("\n"):
        paragraph = copy.deepcopy(template)
        runs = paragraph.findall(qn("a:r"))
        for run in runs[1:]:
            paragraph.remove(run)
        run = runs[0]
        for text_element in run.findall(qn("a:t")):
            run.remove(text_element)
        text_element = run.makeelement(qn("a:t"), {})
        text_element.text = line
        if line != line.strip():
            text_element.set(
                "{http://www.w3.org/XML/1998/namespace}space", "preserve")
        run.append(text_element)
        body.append(paragraph)


def iter_text_frames(slide):
    for node in walk_shapes(slide.shapes):
        shape = node.shape
        if shape.has_text_frame:
            yield shape.text_frame
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text_frame
    if slide.has_notes_slide:
        yield slide.notes_slide.notes_text_frame


def replace_text(prs, old, new):
    return sum(
        replace_in_text_frame(text_frame, old, new)
        for slide in prs.slides
        for text_frame in iter_text_frames(slide)
    )
