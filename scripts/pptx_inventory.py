"""Shared PowerPoint inventory, geometry, and shape locator helpers."""

import hashlib
import os
from collections import Counter
from dataclasses import dataclass

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

EMU_PER_INCH = 914400


class InventoryError(Exception):
    """A presentation cannot be inventoried or targeted safely."""


@dataclass(frozen=True)
class ShapeNode:
    shape: object
    path: tuple
    group_path: tuple
    local_bounds: tuple
    slide_bounds: tuple
    bounds_accuracy: str


def enum_name(value):
    name = getattr(value, "name", None)
    if name:
        return name
    return str(value).split(" (", 1)[0]


def is_picture_shape(shape):
    """Include filled picture placeholders, whose API type is PLACEHOLDER."""
    return shape._element.tag == qn("p:pic")


def sha256_bytes(data):
    return hashlib.sha256(data).hexdigest()


def sha256_file(path):
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _bounds(shape):
    return tuple(int(getattr(shape, key)) for key in
                 ("left", "top", "width", "height"))


def _apply_rect(transform, rect):
    sx, sy, tx, ty = transform
    left, top, width, height = rect
    return (
        int(round(tx + sx * left)),
        int(round(ty + sy * top)),
        int(round(sx * width)),
        int(round(sy * height)),
    )


def _group_child_transform(parent_transform, group):
    xfrm = group._element.grpSpPr.xfrm
    child_width = int(xfrm.chExt.cx)
    child_height = int(xfrm.chExt.cy)
    if child_width == 0 or child_height == 0:
        raise InventoryError(
            f"group {group.name!r} has an invalid child coordinate space")
    parent_sx, parent_sy, parent_tx, parent_ty = parent_transform
    scale_x = int(xfrm.ext.cx) / child_width
    scale_y = int(xfrm.ext.cy) / child_height
    return (
        parent_sx * scale_x,
        parent_sy * scale_y,
        parent_tx + parent_sx * (
            int(xfrm.off.x) - int(xfrm.chOff.x) * scale_x),
        parent_ty + parent_sy * (
            int(xfrm.off.y) - int(xfrm.chOff.y) * scale_y),
    )


def _xml_bool(value):
    return str(value).lower() in {"1", "true"}


def _shape_transform_flags(shape):
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        xfrm = shape._element.grpSpPr.xfrm
    else:
        properties = getattr(shape._element, "spPr", None)
        xfrm = getattr(properties, "xfrm", None)
    if xfrm is None:
        return {"rotation_degrees": 0.0, "flip_h": False, "flip_v": False}
    raw_rotation = xfrm.get("rot")
    return {
        "rotation_degrees": round(int(raw_rotation or 0) / 60000, 3),
        "flip_h": _xml_bool(xfrm.get("flipH")),
        "flip_v": _xml_bool(xfrm.get("flipV")),
    }


def walk_shapes(shapes, transform=(1.0, 1.0, 0.0, 0.0),
                path_prefix=(), group_path=(), bounds_accuracy="exact"):
    """Yield every shape in deterministic pre-order, including groups."""
    for position, shape in enumerate(shapes):
        path = path_prefix + (position,)
        local = _bounds(shape)
        flags = _shape_transform_flags(shape)
        node_accuracy = (
            "approximate" if bounds_accuracy == "approximate"
            or flags["rotation_degrees"] != 0
            or flags["flip_h"] or flags["flip_v"] else "exact"
        )
        node = ShapeNode(
            shape=shape,
            path=path,
            group_path=group_path,
            local_bounds=local,
            slide_bounds=_apply_rect(transform, local),
            bounds_accuracy=node_accuracy,
        )
        yield node
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            parent = {
                "shape_id": int(shape.shape_id),
                "name": shape.name,
                "path": list(path),
            }
            yield from walk_shapes(
                shape.shapes,
                transform=_group_child_transform(transform, shape),
                path_prefix=path,
                group_path=group_path + (parent,),
                bounds_accuracy=node_accuracy,
            )


def _path_node(slide, path):
    if not isinstance(path, (list, tuple)) or not path:
        raise InventoryError("shape path must be a non-empty array")
    shapes = slide.shapes
    shape = None
    for depth, raw_position in enumerate(path):
        try:
            if isinstance(raw_position, bool) or not isinstance(raw_position, int):
                raise TypeError
            position = raw_position
            if position < 0:
                raise IndexError
            shape = shapes[position]
        except (TypeError, ValueError, IndexError) as exc:
            raise InventoryError(f"shape path does not exist: {path!r}") from exc
        if depth < len(path) - 1:
            if shape.shape_type != MSO_SHAPE_TYPE.GROUP:
                raise InventoryError(
                    f"shape path enters non-group shape at {path[:depth + 1]!r}")
            shapes = shape.shapes
    return shape


def recommended_ref(node):
    shape = node.shape
    if getattr(shape, "is_placeholder", False):
        return {"placeholder": int(shape.placeholder_format.idx)}
    return {
        "shape_id": int(shape.shape_id),
        "path": list(node.path),
        "name": shape.name,
    }


def node_for_shape(slide, target):
    for node in walk_shapes(slide.shapes):
        if node.shape._element is target._element:
            return node
    raise InventoryError(f"shape {getattr(target, 'name', None)!r} not found")


def resolve_shape_ref(slide, ref):
    """Resolve a strict inventory locator and cross-check all supplied fields."""
    if not isinstance(ref, dict):
        raise InventoryError("shape reference must be an object")
    nodes = list(walk_shapes(slide.shapes))
    allowed = {"placeholder", "shape_id", "path", "name"}
    unknown = sorted(set(ref) - allowed)
    if unknown:
        raise InventoryError(f"shape reference has unknown fields: {unknown}")
    supplied = allowed & set(ref)
    if not supplied:
        raise InventoryError(
            "shape reference requires placeholder, shape_id, path, or name")

    if "placeholder" in ref:
        placeholder = ref["placeholder"]
        if isinstance(placeholder, bool) or not isinstance(placeholder, int):
            raise InventoryError("placeholder must be an integer")
    if "shape_id" in ref:
        shape_id = ref["shape_id"]
        if isinstance(shape_id, bool) or not isinstance(shape_id, int):
            raise InventoryError("shape_id must be an integer")
    if "name" in ref and not isinstance(ref["name"], str):
        raise InventoryError("name must be a string")
    path_target = _path_node(slide, ref["path"]) if "path" in ref else None

    if "placeholder" in ref:
        matches = [
            node for node in nodes
            if getattr(node.shape, "is_placeholder", False)
            and int(node.shape.placeholder_format.idx) == ref["placeholder"]
        ]
    elif "shape_id" in ref:
        matches = [
            node for node in nodes
            if int(node.shape.shape_id) == ref["shape_id"]
        ]
    elif "path" in ref:
        matches = [node for node in nodes
                   if node.shape._element is path_target._element]
    else:
        matches = [node for node in nodes if node.shape.name == ref["name"]]

    if not matches:
        raise InventoryError(f"shape reference did not match: {ref!r}")
    if len(matches) != 1:
        raise InventoryError(
            f"shape reference is ambiguous ({len(matches)} matches): {ref!r}")
    node = matches[0]
    shape = node.shape

    if "shape_id" in ref and int(ref["shape_id"]) != int(shape.shape_id):
        raise InventoryError(
            f"shape reference ID mismatch: expected {ref['shape_id']}, "
            f"found {shape.shape_id}")
    if "path" in ref and list(ref["path"]) != list(node.path):
        raise InventoryError(
            f"shape reference path mismatch: expected {ref['path']!r}, "
            f"found {list(node.path)!r}")
    if "name" in ref and ref["name"] != shape.name:
        raise InventoryError(
            f"shape reference name mismatch: expected {ref['name']!r}, "
            f"found {shape.name!r}")
    if "placeholder" in ref:
        actual = (int(shape.placeholder_format.idx)
                  if getattr(shape, "is_placeholder", False) else None)
        if int(ref["placeholder"]) != actual:
            raise InventoryError(
                f"placeholder mismatch: expected {ref['placeholder']}, "
                f"found {actual}")
    return shape


def _rect_record(rect):
    left, top, width, height = rect
    return {"left": left, "top": top, "width": width, "height": height}


def _inches_record(rect):
    return {
        key: round(value / EMU_PER_INCH, 3)
        for key, value in _rect_record(rect).items()
    }


def _font_color(font):
    try:
        rgb = font.color.rgb
        if rgb is not None:
            return str(rgb)
    except (AttributeError, ValueError):
        pass
    try:
        theme = font.color.theme_color
        if theme is not None:
            return f"theme:{enum_name(theme)}"
    except (AttributeError, ValueError):
        pass
    return None


def text_style_summary(text_frame):
    sizes = set()
    fonts = set()
    colors = set()
    for paragraph in text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.size is not None:
                sizes.add(round(run.font.size.pt, 3))
            if run.font.name:
                fonts.add(run.font.name)
            color = _font_color(run.font)
            if color:
                colors.add(color)
    return {
        "max_font_size_pt": max(sizes) if sizes else None,
        "font_sizes_pt": sorted(sizes),
        "fonts": sorted(fonts),
        "colors": sorted(colors),
    }


def _color_format_value(color_format):
    try:
        rgb = color_format.rgb
        if rgb is not None:
            return str(rgb)
    except (AttributeError, ValueError):
        pass
    try:
        theme = color_format.theme_color
        if theme is not None:
            return f"theme:{enum_name(theme)}"
    except (AttributeError, ValueError):
        pass
    return None


def _fill_record(shape):
    try:
        fill = shape.fill
        fill_type = enum_name(fill.type) if fill.type is not None else None
    except (AttributeError, TypeError, ValueError):
        return None
    result = {"type": fill_type}
    try:
        color = _color_format_value(fill.fore_color)
    except (AttributeError, TypeError, ValueError):
        color = None
    if color is not None:
        result["color"] = color
    try:
        transparency = fill.transparency
    except (AttributeError, TypeError, ValueError):
        transparency = None
    if transparency is not None:
        result["transparency"] = transparency
    return result


def _line_record(shape):
    try:
        line = shape.line
    except (AttributeError, TypeError, ValueError):
        return None
    result = {}
    try:
        color = _color_format_value(line.color)
    except (AttributeError, TypeError, ValueError):
        color = None
    if color is not None:
        result["color"] = color
    try:
        result["width_emu"] = int(line.width) if line.width is not None else None
    except (AttributeError, TypeError, ValueError):
        result["width_emu"] = None
    try:
        result["dash_style"] = (
            enum_name(line.dash_style) if line.dash_style is not None else None)
    except (AttributeError, TypeError, ValueError):
        result["dash_style"] = None
    return result


def _paragraph_record(paragraph):
    p_pr = paragraph._p.find(qn("a:pPr"))
    bullet = False
    if p_pr is not None:
        bullet = any(
            child.tag in {qn("a:buChar"), qn("a:buAutoNum"), qn("a:buBlip")}
            for child in p_pr
        )
    runs = []
    for run in paragraph.runs:
        runs.append({
            "text": run.text,
            "font": run.font.name,
            "size_pt": (round(run.font.size.pt, 3)
                        if run.font.size is not None else None),
            "bold": run.font.bold,
            "italic": run.font.italic,
            "color": _font_color(run.font),
            "hyperlink": run.hyperlink.address,
        })
    return {
        "text": paragraph.text,
        "level": int(paragraph.level),
        "alignment": (enum_name(paragraph.alignment)
                      if paragraph.alignment is not None else None),
        "bullet": bullet,
        "runs": runs,
    }


def _text_frame_record(text_frame):
    def optional_emu(value):
        return int(value) if value is not None else None

    return {
        "margin_emu": {
            "left": optional_emu(text_frame.margin_left),
            "right": optional_emu(text_frame.margin_right),
            "top": optional_emu(text_frame.margin_top),
            "bottom": optional_emu(text_frame.margin_bottom),
        },
        "word_wrap": text_frame.word_wrap,
        "auto_size": (enum_name(text_frame.auto_size)
                      if text_frame.auto_size is not None else None),
        "vertical_anchor": (enum_name(text_frame.vertical_anchor)
                            if text_frame.vertical_anchor is not None else None),
    }


def chart_record(chart):
    plots = list(chart.plots)
    categories = []
    if plots:
        try:
            categories = [str(value) for value in plots[0].categories]
        except (AttributeError, TypeError, ValueError):
            categories = []
    series_records = []
    for plot in plots:
        for series in plot.series:
            try:
                name = series.name
            except (AttributeError, KeyError):
                name = None
            try:
                values = list(series.values)
            except (AttributeError, TypeError, ValueError):
                values = []
            series_records.append({"name": name, "values": values})
    title = None
    if chart.has_title:
        title = chart.chart_title.text_frame.text
    return {
        "chart_type": enum_name(chart.chart_type),
        "title": title,
        "categories": categories,
        "series": series_records,
    }


def _placeholder_record(shape):
    if not getattr(shape, "is_placeholder", False):
        return None
    return {
        "idx": int(shape.placeholder_format.idx),
        "type": enum_name(shape.placeholder_format.type),
    }


def shape_record(node, warnings):
    shape = node.shape
    is_picture = is_picture_shape(shape)
    is_group = shape.shape_type == MSO_SHAPE_TYPE.GROUP
    capabilities = {
        "text": bool(shape.has_text_frame),
        "picture": is_picture,
        "table": bool(shape.has_table),
        "chart": bool(shape.has_chart),
        "group": is_group,
    }
    content = {}
    transform_flags = _shape_transform_flags(shape)
    try:
        auto_shape_type = (enum_name(shape.auto_shape_type)
                           if shape.auto_shape_type is not None else None)
    except (AttributeError, TypeError, ValueError):
        auto_shape_type = None
    record = {
        "shape_id": int(shape.shape_id),
        "name": shape.name,
        "type": enum_name(shape.shape_type),
        "auto_shape_type": auto_shape_type,
        "path": list(node.path),
        "z_order": int(node.path[-1]),
        "group_path": list(node.group_path),
        "recommended_ref": recommended_ref(node),
        "placeholder": _placeholder_record(shape),
        "bounds": {
            "accuracy": node.bounds_accuracy,
            "local_emu": _rect_record(node.local_bounds),
            "slide_emu": _rect_record(node.slide_bounds),
            "slide_inches": _inches_record(node.slide_bounds),
        },
        "transform": transform_flags,
        "style": {
            "fill": _fill_record(shape),
            "line": _line_record(shape),
        },
        "capabilities": capabilities,
        "content": content,
    }
    if node.bounds_accuracy == "approximate":
        warnings.append({
            "code": "approximate_bounds",
            "message": "rotation or flip is not included in absolute bounds",
            "shape_id": int(shape.shape_id),
            "path": list(node.path),
        })

    if shape.has_text_frame:
        content["text"] = shape.text_frame.text
        content["paragraph_count"] = len(shape.text_frame.paragraphs)
        content["paragraphs"] = [
            _paragraph_record(paragraph)
            for paragraph in shape.text_frame.paragraphs
        ]
        record["text_style"] = text_style_summary(shape.text_frame)
        record["text_frame"] = _text_frame_record(shape.text_frame)

    if shape.has_table:
        table = shape.table
        content["table"] = {
            "row_count": len(table.rows),
            "column_count": len(table.columns),
            "cells": [
                [cell.text for cell in row.cells] for row in table.rows
            ],
        }

    if shape.has_chart:
        content["chart"] = chart_record(shape.chart)

    if is_picture:
        try:
            image = shape.image
            blob = image.blob
            content["image"] = {
                "filename": image.filename,
                "extension": image.ext,
                "content_type": getattr(image, "content_type", None),
                "size_bytes": len(blob),
                "sha256": sha256_bytes(blob),
            }
        except (KeyError, ValueError) as exc:
            content["image"] = {"readable": False}
            warnings.append({
                "code": "unreadable_image",
                "message": str(exc),
                "shape_id": int(shape.shape_id),
            })

    if is_group:
        content["group"] = {
            "children": [int(child.shape_id) for child in shape.shapes]
        }

    if not any(capabilities.values()):
        warnings.append({
            "code": "unmodeled_shape",
            "message": f"shape type {record['type']} has no modeled content",
            "shape_id": int(shape.shape_id),
            "path": list(node.path),
        })
    return record


def _validate_slide_identity(nodes, slide_index):
    ids = [int(node.shape.shape_id) for node in nodes]
    duplicates = [value for value, count in Counter(ids).items() if count > 1]
    if duplicates:
        raise InventoryError(
            f"slide {slide_index} contains duplicate shape IDs: {duplicates}")


def slide_inventory(slide, slide_index, layout_index):
    nodes = list(walk_shapes(slide.shapes))
    _validate_slide_identity(nodes, slide_index)
    warnings = []
    names = Counter(node.shape.name for node in nodes)
    for name, count in sorted(names.items()):
        if count > 1:
            warnings.append({
                "code": "duplicate_shape_name",
                "message": f"shape name {name!r} occurs {count} times",
                "name": name,
                "shape_ids": [
                    int(node.shape.shape_id) for node in nodes
                    if node.shape.name == name
                ],
            })
    placeholder_indexes = [
        int(node.shape.placeholder_format.idx) for node in nodes
        if getattr(node.shape, "is_placeholder", False)
    ]
    for value, count in sorted(Counter(placeholder_indexes).items()):
        if count > 1:
            warnings.append({
                "code": "duplicate_placeholder",
                "message": f"placeholder {value} occurs {count} times",
                "placeholder": value,
            })

    shapes = [shape_record(node, warnings) for node in nodes]
    notes = (slide.notes_slide.notes_text_frame.text
             if slide.has_notes_slide else None)
    return {
        "index": slide_index,
        "slide_id": int(slide.slide_id),
        "recommended_ref": {
            "slide_id": int(slide.slide_id), "index": slide_index,
        },
        "layout": {"index": layout_index, "name": slide.slide_layout.name},
        "notes": notes,
        "shapes": shapes,
        "warnings": warnings,
    }


def _inherited_shape_inventory(shapes, scope):
    nodes = list(walk_shapes(shapes))
    _validate_slide_identity(nodes, scope)
    warnings = []
    records = []
    for node in nodes:
        record = shape_record(node, warnings)
        record.pop("recommended_ref", None)
        record["editable_from_slide"] = False
        records.append(record)
    return records, warnings


def _layout_inventory(layout, index):
    placeholders = []
    for shape in layout.placeholders:
        placeholders.append({
            "idx": int(shape.placeholder_format.idx),
            "type": enum_name(shape.placeholder_format.type),
            "name": shape.name,
            "bounds_emu": _rect_record(_bounds(shape)),
            "bounds_inches": _inches_record(_bounds(shape)),
        })
    shapes, warnings = _inherited_shape_inventory(
        layout.shapes, f"layout {index}")
    return {
        "index": index,
        "recommended_ref": {"index": index, "name": layout.name},
        "name": layout.name,
        "placeholders": placeholders,
        "shapes": shapes,
        "warnings": warnings,
    }


def _master_inventory(master, index):
    shapes, warnings = _inherited_shape_inventory(
        master.shapes, f"master {index}")
    return {
        "index": index,
        "shapes": shapes,
        "warnings": warnings,
    }


def inspect_presentation(path):
    prs = Presentation(path)
    layouts = list(prs.slide_layouts)
    layout_indexes = {id(layout._element): index
                      for index, layout in enumerate(layouts)}
    slides = []
    for index, slide in enumerate(prs.slides):
        layout_index = layout_indexes.get(id(slide.slide_layout._element))
        slides.append(slide_inventory(slide, index, layout_index))
    layout_records = [
        _layout_inventory(layout, index)
        for index, layout in enumerate(layouts)
    ]
    master_records = [
        _master_inventory(master, index)
        for index, master in enumerate(prs.slide_masters)
    ]
    warnings = [
        dict(item, scope="slide", slide_index=slide["index"])
        for slide in slides for item in slide["warnings"]
    ]
    warnings.extend(
        dict(item, scope="layout", layout_index=layout["index"])
        for layout in layout_records for item in layout["warnings"]
    )
    warnings.extend(
        dict(item, scope="master", master_index=master["index"])
        for master in master_records for item in master["warnings"]
    )
    return {
        "schema": "powerpoint.inventory/v2",
        "ok": True,
        "template": {
            "path": os.path.abspath(path),
            "filename": os.path.basename(path),
            "sha256": sha256_file(path),
            "size_bytes": os.path.getsize(path),
        },
        "presentation": {
            "slide_count": len(prs.slides),
            "layout_count": len(layouts),
            "master_count": len(prs.slide_masters),
            "slide_size": {
                "width_emu": int(prs.slide_width),
                "height_emu": int(prs.slide_height),
                "width_inches": round(prs.slide_width / EMU_PER_INCH, 3),
                "height_inches": round(prs.slide_height / EMU_PER_INCH, 3),
            },
        },
        "masters": master_records,
        "layouts": layout_records,
        "slides": slides,
        "warnings": warnings,
    }


def export_images(path, out_dir):
    prs = Presentation(path)
    os.makedirs(out_dir, exist_ok=True)
    exports = []
    for slide_index, slide in enumerate(prs.slides):
        nodes = list(walk_shapes(slide.shapes))
        _validate_slide_identity(nodes, slide_index)
        for node in nodes:
            shape = node.shape
            if not is_picture_shape(shape):
                continue
            try:
                image = shape.image
            except (KeyError, ValueError):
                continue
            output = os.path.join(
                out_dir,
                f"slide-{slide_index:03d}-shape-{int(shape.shape_id):04d}."
                f"{image.ext}",
            )
            with open(output, "wb") as handle:
                handle.write(image.blob)
            exports.append({
                "slide_index": slide_index,
                "shape_ref": recommended_ref(node),
                "output": output,
                "sha256": sha256_bytes(image.blob),
                "size_bytes": len(image.blob),
            })
    return {
        "schema": "powerpoint.image-export/v1",
        "ok": True,
        "template": {
            "path": os.path.abspath(path),
            "sha256": sha256_file(path),
        },
        "exports": exports,
    }
