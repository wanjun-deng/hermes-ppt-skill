#!/usr/bin/env python3
"""Read-only environment and optional rendering probe for powerpoint-dev."""
import argparse
import importlib.util
import json
from pathlib import Path
import shutil
import sys

from pptx import Presentation
from pptx_common import atomic_save_presentation
from pptx_render import render
from pptx_fonts import probe_font


def doctor(probe_dir=None, fonts=()):
    modules = {name: importlib.util.find_spec(name) is not None
               for name in ("pptx", "lxml", "PIL", "jsonschema", "docx", "pypdf", "openpyxl")}
    result = {"schema": "powerpoint.environment/v1", "python": sys.executable,
              "modules": modules, "tools": {name: shutil.which(name) for name in
              ("soffice", "pdftoppm", "pdftocairo", "fc-match")},
              "render_probe": {"status": "not_run"}}
    if probe_dir:
        folder = Path(probe_dir).resolve()
        folder.mkdir(parents=True, exist_ok=True)
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[0])
        slide.shapes.title.text = "视觉连通测试 VISUAL 427"
        slide.placeholders[1].text = "蓝色方块 / Blue square"
        from pptx.util import Inches
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.dml.color import RGBColor
        square = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4), Inches(4), Inches(1), Inches(1))
        square.fill.solid()
        square.fill.fore_color.rgb = RGBColor(0, 80, 220)
        source = atomic_save_presentation(deck, folder / "probe.pptx")
        try:
            manifest = render(source, str(folder / "render"), "slide", 120)
            result["render_probe"] = {"status": "passed" if manifest["rendered"] else "unavailable", "manifest": manifest}
        except Exception as exc:
            result["render_probe"] = {"status": "failed", "error": str(exc)}
    result["ok"] = all(modules[x] for x in ("pptx", "lxml", "PIL", "jsonschema")) and result["render_probe"]["status"] not in {"failed", "unavailable"}
    result["fonts"] = [probe_font(family) for family in fonts]
    if fonts:
        result["ok"] = result["ok"] and all(row["status"] == "matched" and row["glyphs_available"] for row in result["fonts"])
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--probe-dir")
    parser.add_argument("--output")
    parser.add_argument("--font", action="append", default=[], help="probe a requested Chinese font (repeatable)")
    args = parser.parse_args()
    report = doctor(args.probe_dir, args.font)
    if args.output:
        Path(args.output).write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    raise SystemExit(0 if report["ok"] else 1)
