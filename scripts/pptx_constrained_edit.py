#!/usr/bin/env python3
"""Execute a hash-bound, narrowly scoped edit and prove its package changes."""
import argparse
import copy
import hashlib
import json
from pathlib import Path
import posixpath
import sys
import zipfile

from jsonschema import Draft202012Validator
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from pptx_common import atomic_save_presentation, write_error
from pptx_fonts import set_shape_fonts, text_frames
from pptx_geometry import geometry_report
from pptx_inventory import resolve_shape_ref, sha256_file
from pptx_text import replace_in_text_frame

ROOT = Path(__file__).resolve().parents[1]


def canonical(element):
    """Ignore serialization indentation, retaining significant text and attrs."""
    element = copy.deepcopy(element)
    for node in element.iter():
        if len(node) and node.text and not node.text.strip():
            node.text = None
        if node.tail and not node.tail.strip():
            node.tail = None
    return etree.tostring(element, method="c14n")


def package(path):
    with zipfile.ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist() if not name.endswith("/")}


def comparable(name, data):
    return canonical(etree.fromstring(data)) if name.endswith((".xml", ".rels")) else data


def _mask_target(slide_xml, shape_id):
    root = copy.deepcopy(slide_xml)
    for node in root.iter(qn("p:cNvPr")):
        if node.get("id") == str(shape_id):
            shape = node.getparent().getparent()
            shape.getparent().replace(shape, etree.Element("scope-target"))
            return canonical(root)
    raise ValueError("target disappeared from slide")


def _projection(shape_xml, operation):
    root = copy.deepcopy(shape_xml)
    action = operation["action"]
    if action == "replace_text":
        for node in root.iter(qn("a:t")):
            node.text = ""
    elif action in {"set_font", "set_text_style"}:
        for properties in list(root.iter(qn("a:rPr"))):
            if action == "set_font":
                for slot, tag in (("latin", "latin"), ("east_asian", "ea"), ("complex_script", "cs")):
                    if slot in operation:
                        for child in list(properties):
                            if child.tag == qn("a:" + tag):
                                properties.remove(child)
            else:
                for field, attr in (("font_size_pt", "sz"), ("bold", "b")):
                    if field in operation:
                        properties.attrib.pop(attr, None)
                if "color" in operation:
                    for child in list(properties):
                        if child.tag == qn("a:solidFill"):
                            properties.remove(child)
            if not len(properties) and not properties.attrib:
                properties.getparent().remove(properties)
        if "alignment" in operation:
            for properties in list(root.iter(qn("a:pPr"))):
                properties.attrib.pop("algn", None)
                if not len(properties) and not properties.attrib:
                    properties.getparent().remove(properties)
    elif action == "set_geometry":
        properties = root.find(qn("p:spPr"))
        transform = properties.find(qn("a:xfrm")) if properties is not None else root.find(qn("p:xfrm"))
        if transform is not None:
            for tag in ("off", "ext"):
                node = transform.find(qn("a:" + tag))
                if node is not None:
                    transform.remove(node)
            if not len(transform) and not transform.attrib:
                transform.getparent().remove(transform)
        if properties is not None and not len(properties) and not properties.attrib:
            root.remove(properties)
    return canonical(root)


def _style(shape, operation):
    for frame in text_frames(shape):
        for paragraph in frame.paragraphs:
            if "alignment" in operation:
                paragraph.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[operation["alignment"]]
            for run in paragraph.runs:
                if "font_size_pt" in operation:
                    run.font.size = Pt(operation["font_size_pt"])
                if "bold" in operation:
                    run.font.bold = operation["bold"]
                if "color" in operation:
                    run.font.color.rgb = RGBColor.from_string(operation["color"])


def _add_text(slide, item):
    box = item["bounds"]
    shape = slide.shapes.add_textbox(*(Inches(box[k]) for k in ("left", "top", "width", "height")))
    shape.text = item["text"]
    if item.get("style"):
        _style(shape, item["style"])
        set_shape_fonts(shape, **{k: item["style"][k] for k in ("latin", "east_asian") if k in item["style"]})
    return shape


def _slide(prs, ref):
    index = ref["index"]
    if index >= len(prs.slides) or prs.slides[index].slide_id != ref["slide_id"]:
        raise ValueError("slide locator does not match the source")
    return prs.slides[index]


def _assert_package(before, after, allowed, appended):
    changed = [name for name in before.keys() & after.keys()
               if comparable(name, before[name]) != comparable(name, after[name])]
    removed = sorted(before.keys() - after.keys())
    added = sorted(after.keys() - before.keys())
    if removed or set(changed) - allowed:
        raise ValueError("out-of-scope package changes: " + str(removed + sorted(set(changed) - allowed)))
    permitted_new = set()
    for part in appended:
        permitted_new.update((part, posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part)+".rels")))
    if set(added) - permitted_new:
        raise ValueError("unexpected new package parts: " + str(sorted(set(added) - permitted_new)))
    if appended:
        # Preserve all original relationships, content types, and presentation
        # properties; only explicitly appended slide references may be added.
        for part in ("ppt/presentation.xml", "ppt/_rels/presentation.xml.rels", "[Content_Types].xml"):
            old, new = etree.fromstring(before[part]), etree.fromstring(after[part])
            if part == "ppt/presentation.xml":
                old = old.find(qn("p:sldIdLst"))
                new = new.find(qn("p:sldIdLst"))
                originals = list(old) if old is not None else []
                if new is None or [canonical(e) for e in list(new)[:len(originals)]] != [canonical(e) for e in originals]:
                    raise ValueError("original slide order changed")
                old_root, new_root = etree.fromstring(before[part]), etree.fromstring(after[part])
                for root in (old_root, new_root):
                    lst = root.find(qn("p:sldIdLst"))
                    if lst is not None:
                        root.remove(lst)
                if canonical(old_root) != canonical(new_root):
                    raise ValueError("presentation properties changed")
            else:
                old_entries = {e.get("Id") or e.get("PartName") or e.get("Extension"): canonical(e) for e in old}
                new_entries = {e.get("Id") or e.get("PartName") or e.get("Extension"): canonical(e) for e in new}
                if any(new_entries.get(k) != value for k, value in old_entries.items()):
                    raise ValueError("existing package relationship or content type changed")
                extra = [e for e in new if (e.get("Id") or e.get("PartName") or e.get("Extension")) not in old_entries]
                if len(extra) != len(appended):
                    raise ValueError("unexpected appended package reference count")
                if part.endswith(".rels"):
                    targets = set()
                    for entry in extra:
                        target = posixpath.normpath(posixpath.join("ppt", entry.get("Target", ""))).lstrip("/")
                        if (entry.get("Type") != "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide"
                                or entry.get("TargetMode") or target not in appended):
                            raise ValueError("unexpected appended relationship")
                        targets.add(target)
                    if targets != set(appended):
                        raise ValueError("appended slide relationship mismatch")
                elif {e.get("PartName", "").lstrip("/") for e in extra} != set(appended) or any(
                    e.get("ContentType") != "application/vnd.openxmlformats-officedocument.presentationml.slide+xml" for e in extra
                ):
                    raise ValueError("unexpected appended content type")
    return {"changed": sorted(changed), "added": added, "removed": removed}


def execute_edit(source, spec, output):
    schema = json.loads((ROOT / "schemas/powerpoint-edit-v1.schema.json").read_text(encoding="utf-8"))
    json.dumps(spec, allow_nan=False)
    Draft202012Validator(schema).validate(spec)
    source, output = Path(source).resolve(), Path(output).resolve()
    if source == output:
        raise ValueError("constrained edits require a separate output; preserve the source baseline")
    source_hash = sha256_file(source)
    if source_hash != spec["source"]["sha256"]:
        raise ValueError("source hash mismatch")
    prs = Presentation(source)
    if len(prs.slides) != spec["source"]["slide_count"]:
        raise ValueError("source slide count mismatch")
    original = package(source)
    allowed, appended, audit = set(), [], []
    mode_actions = {"restyle": {"set_font", "set_text_style", "set_geometry"},
                    "precise": {"replace_text", "set_font", "set_text_style", "set_geometry"},
                    "enrich": {"add_text", "append_slide"}}
    for index, operation in enumerate(spec["operations"]):
        action = operation["action"]
        if action not in mode_actions[spec["mode"]]:
            raise ValueError(f"{action} is forbidden in {spec['mode']} mode")
        if action == "append_slide":
            ref = operation["layout"]
            if ref["index"] >= len(prs.slide_layouts) or prs.slide_layouts[ref["index"]].name != ref["name"]:
                raise ValueError("layout locator mismatch")
            slide = prs.slides.add_slide(prs.slide_layouts[ref["index"]])
            for item in operation["text"]:
                _add_text(slide, item)
            part = str(slide.part.partname).lstrip("/")
            appended.append(part)
            allowed.update(("ppt/presentation.xml", "ppt/_rels/presentation.xml.rels", "[Content_Types].xml"))
        else:
            slide = _slide(prs, operation["slide"])
            before_slide = copy.deepcopy(slide._element)
            if action == "add_text":
                shape = _add_text(slide, operation)
                restored = copy.deepcopy(slide._element)
                for node in restored.iter(qn("p:cNvPr")):
                    if node.get("id") == str(shape.shape_id):
                        element = node.getparent().getparent()
                        element.getparent().remove(element)
                        break
                if canonical(restored) != canonical(before_slide):
                    raise ValueError("addition modified existing slide objects")
            else:
                shape = resolve_shape_ref(slide, operation["shape"])
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    raise ValueError("target a group's leaf object explicitly")
                before_shape = copy.deepcopy(shape._element)
                frames = list(text_frames(shape))
                if action in {"replace_text", "set_font", "set_text_style"} and not frames:
                    raise ValueError("target has no editable text")
                if action == "replace_text":
                    expected = [frame.text.replace(operation["old"], operation["new"]) for frame in frames]
                    hits = sum(replace_in_text_frame(frame, operation["old"], operation["new"]) for frame in frames)
                    if hits != operation.get("expected_count", 1) or [frame.text for frame in frames] != expected:
                        raise ValueError("text occurrence count or replacement result mismatch")
                elif action == "set_font":
                    set_shape_fonts(shape, **{k: operation[k] for k in ("latin", "east_asian", "complex_script") if k in operation})
                elif action == "set_text_style":
                    _style(shape, operation)
                elif action == "set_geometry":
                    for key, value in operation["bounds"].items():
                        setattr(shape, key, Inches(value))
                if _projection(before_shape, operation) != _projection(shape._element, operation):
                    raise ValueError("operation changed a frozen target property")
                if _mask_target(before_slide, shape.shape_id) != _mask_target(slide._element, shape.shape_id):
                    raise ValueError("operation changed a non-target object")
            part = str(slide.part.partname).lstrip("/")
        allowed.add(part)
        audit.append({"operation_index": index, "action": action, "slide_id": slide.slide_id})
    expected_ids = [s.slide_id for s in prs.slides]
    expected_slides = {str(s.part.partname).lstrip("/"): canonical(s._element) for s in prs.slides}
    verified, verified_geometry = {}, {}
    def validate(candidate):
        if sha256_file(source) != source_hash:
            raise ValueError("source changed during the edit")
        reopened = Presentation(candidate)
        if [s.slide_id for s in reopened.slides] != expected_ids:
            raise ValueError("source slide identity or order changed")
        for slide in reopened.slides:
            if canonical(slide._element) != expected_slides[str(slide.part.partname).lstrip("/")]:
                raise ValueError("saved slide differs from the verified edit")
        verified.update(_assert_package(original, package(candidate), allowed, appended))
        verified_geometry.update(geometry_report(candidate, source))
    atomic_save_presentation(prs, output, validator=validate)
    return {"schema": "powerpoint.edit-proof/v1", "source_sha256": source_hash,
            "artifact_sha256": sha256_file(output), "output": str(output),
            "scope_verified": True, "visual_review_pending": True, "mode": spec["mode"],
            "contract": spec, "operations": audit, "package_changes": verified,
            "source_parts": {name: hashlib.sha256(data).hexdigest() for name, data in original.items()},
            "geometry": verified_geometry}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source")
    parser.add_argument("contract")
    parser.add_argument("output")
    parser.add_argument("--report", required=True)
    args = parser.parse_args(argv)
    try:
        if Path(args.report).resolve() in {Path(v).resolve() for v in (args.source, args.contract, args.output)}:
            raise ValueError("report must not overwrite an input or the PPTX")
        spec = json.loads(Path(args.contract).read_text(encoding="utf-8"))
        result = execute_edit(args.source, spec, args.output)
        from pptx_workflow import save_json
        save_json(args.report, result)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    except Exception as exc:
        write_error(exc, stream=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
