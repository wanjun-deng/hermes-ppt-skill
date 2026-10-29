#!/usr/bin/env python3
"""Extract source-addressable text, tables and pictures for presentation planning."""
import argparse
import csv
from datetime import date, datetime
import json
from pathlib import Path
import sys

from PIL import Image
from jsonschema import Draft202012Validator
from pptx_common import write_error
from pptx_inventory import sha256_file


def json_value(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else value


def extract_materials(paths, asset_dir):
    assets = Path(asset_dir).resolve()
    assets.mkdir(parents=True, exist_ok=True)
    files, records, issues = [], [], []
    def record(file_id, kind, location, **data):
        row = {"id": file_id + ":" + str(sum(r["file_id"] == file_id for r in records) + 1),
               "file_id": file_id, "kind": kind, "location": location, **data}
        records.append(row)
        return row
    for file_index, value in enumerate(paths):
        path = Path(value).expanduser().resolve()
        file_id = "file" + str(file_index + 1)
        files.append({"id": file_id, "path": str(path), "sha256": sha256_file(path)})
        suffix = path.suffix.lower()
        if suffix in {".txt", ".md"}:
            for number, line in enumerate(path.read_text(encoding="utf-8-sig").splitlines(), 1):
                if line.strip():
                    record(file_id, "text", {"line": number}, text=line)
        elif suffix == ".csv":
            with path.open(encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.reader(handle))
            for index, row in enumerate(rows, 1):
                record(file_id, "table_row", {"row": index}, cells=row, text=" | ".join(row))
        elif suffix == ".docx":
            from docx import Document
            from docx.table import Table
            doc = Document(path)
            for index, block in enumerate(doc.iter_inner_content(), 1):
                if isinstance(block, Table):
                    rows = [[cell.text for cell in row.cells] for row in block.rows]
                    record(file_id, "table", {"block": index}, rows=rows, text="\n".join(" | ".join(row) for row in rows))
                elif block.text.strip():
                    record(file_id, "text", {"block": index}, text=block.text)
            for rel_id, rel in doc.part.rels.items():
                if rel.reltype.endswith("/image") and not rel.is_external:
                    part = rel.target_part
                    target = assets / (file_id + "-" + rel_id + Path(str(part.partname)).suffix)
                    target.write_bytes(part.blob)
                    with Image.open(target) as image:
                        record(file_id, "image", {"relationship": rel_id}, path=str(target), width_px=image.width, height_px=image.height, sha256=sha256_file(target))
            issues.append({"file_id": file_id, "severity": "info", "code": "docx_scope", "message": "Extracted body blocks and embedded images; headers, footnotes, revisions and text boxes need separate inspection if required."})
        elif suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(path)
            for number, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                if text.strip():
                    record(file_id, "text", {"page": number}, text=text)
                else:
                    issues.append({"file_id": file_id, "page": number, "severity": "blocking", "code": "no_extractable_pdf_text", "message": "Page has no extractable text; OCR or visual reading is required, not implemented by this reader."})
            issues.append({"file_id": file_id, "severity": "info", "code": "pdf_reading_order", "message": "Page text extraction does not guarantee column order or table structure; check the original page."})
        elif suffix == ".xlsx":
            from openpyxl import load_workbook
            formulas = load_workbook(path, data_only=False, read_only=True)
            cached = load_workbook(path, data_only=True, read_only=True)
            try:
                for sheet in formulas:
                    for row in sheet.iter_rows():
                        for cell in row:
                            if cell.value is None:
                                continue
                            value = cell.value
                            formula = None
                            if cell.data_type == "f":
                                formula = value
                                value = cached[sheet.title][cell.coordinate].value
                                if value is None:
                                    issues.append({"file_id": file_id, "sheet": sheet.title, "cell": cell.coordinate, "severity": "blocking", "code": "formula_cache_missing", "message": "Formula has no cached result; no calculation was performed."})
                            record(file_id, "cell", {"sheet": sheet.title, "cell": cell.coordinate}, value=json_value(value), formula=formula, number_format=cell.number_format, text="" if value is None else str(json_value(value)))
            finally:
                formulas.close()
                cached.close()
        elif suffix in {".png", ".jpg", ".jpeg", ".webp", ".gif", ".tif", ".tiff"}:
            with Image.open(path) as image:
                record(file_id, "image", {"image": 1}, path=str(path), width_px=image.width, height_px=image.height, sha256=sha256_file(path))
        else:
            issues.append({"file_id": file_id, "severity": "blocking", "code": "unsupported_material", "message": "Unsupported file type: " + suffix})
    return {"schema": "powerpoint.materials/v1", "files": files, "records": records, "issues": issues,
            "ready": not any(issue["severity"] == "blocking" for issue in issues)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("paths", nargs="+")
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    try:
        output = Path(args.output).resolve()
        result = extract_materials(args.paths, output.parent / "material-assets")
        schema = Path(__file__).resolve().parents[1] / "schemas/powerpoint-materials-v1.schema.json"
        Draft202012Validator(json.loads(schema.read_text())).validate(result)
        output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(json.dumps({"output": str(output), "ready": result["ready"], "records": len(result["records"]), "issues": result["issues"]}, ensure_ascii=False, indent=2))
        return 0 if result["ready"] else 1
    except Exception as exc:
        write_error(exc, stream=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
