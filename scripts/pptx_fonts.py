#!/usr/bin/env python3
"""Explicit font assignment and read-only fontconfig coverage probes."""
import argparse
import json
import shutil
import subprocess
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement


def set_run_fonts(run, *, latin=None, east_asian=None, complex_script=None):
    """Assign only requested slots, retaining inherited fonts otherwise."""
    if latin is not None:
        run.font.name = latin
    for tag, family, successors in (
        ("ea", east_asian, ("a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")),
        ("cs", complex_script, ("a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")),
    ):
        if family is None:
            continue
        properties = run._r.get_or_add_rPr()
        element = properties.find(qn("a:" + tag))
        if element is None:
            element = OxmlElement("a:" + tag)
            properties.insert_element_before(element, *successors)
        element.set("typeface", family)


def text_frames(shape):
    if shape.has_text_frame:
        yield shape.text_frame
    if shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                yield cell.text_frame


def set_shape_fonts(shape, **fonts):
    for frame in text_frames(shape):
        for paragraph in frame.paragraphs:
            for run in paragraph.runs:
                set_run_fonts(run, **fonts)


def set_presentation_fonts(prs, **fonts):
    """Explicit creation override, including chart defaults and title runs."""
    from pptx_inventory import walk_shapes
    for slide in prs.slides:
        for node in walk_shapes(slide.shapes):
            set_shape_fonts(node.shape, **fonts)
            if node.shape.has_chart:
                chart = node.shape.chart
                if chart.has_title:
                    for paragraph in chart.chart_title.text_frame.paragraphs:
                        for run in paragraph.runs:
                            set_run_fonts(run, **fonts)
                if fonts.get("latin"):
                    chart.font.name = fonts["latin"]
                if fonts.get("east_asian"):
                    properties = chart.font._rPr
                    element = properties.find(qn("a:ea"))
                    if element is None:
                        element = OxmlElement("a:ea")
                        properties.insert_element_before(element, "a:cs", "a:sym", "a:hlinkClick", "a:extLst")
                    element.set("typeface", fonts["east_asian"])


def probe_font(family, text="中文测试"):
    result = {"requested": family, "status": "unconfirmed", "matched_families": [],
              "file": None, "glyphs_available": None, "missing_characters": []}
    tool = shutil.which("fc-match")
    if not tool:
        result["reason"] = "fontconfig unavailable; resolution was not checked"
        return result
    try:
        process = subprocess.run([tool, "-f", "%{family}\n%{file}\n%{charset}\n", family],
                                 capture_output=True, text=True, timeout=10, check=True)
        names, path, charset = process.stdout.splitlines()[:3]
        result.update(matched_families=[v.strip() for v in names.split(",")], file=path)
        ranges = [tuple(int(v, 16) for v in token.split("-")) for token in charset.split()]
        missing = sorted({c for c in text if not c.isspace() and not any(
            row[0] <= ord(c) <= row[-1] for row in ranges)})
        result.update(glyphs_available=not missing, missing_characters=missing)
        result["status"] = "matched" if family.casefold() in {
            name.casefold() for name in result["matched_families"]} else "substituted"
    except (OSError, subprocess.SubprocessError, ValueError, IndexError) as exc:
        result["reason"] = str(exc)
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("families", nargs="+")
    parser.add_argument("--text", default="中文测试")
    args = parser.parse_args()
    rows = [probe_font(family, args.text) for family in args.families]
    print(json.dumps({"schema": "powerpoint.font-probe/v1", "fonts": rows}, ensure_ascii=False, indent=2))
    raise SystemExit(0 if all(r["status"] == "matched" and r["glyphs_available"] for r in rows) else 1)
