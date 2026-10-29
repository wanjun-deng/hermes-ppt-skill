#!/usr/bin/env python3
"""Fill tokens in a template and optionally append layout-based slides."""
import argparse
import json
import re
import sys

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from pptx_common import (PptxSkillError, atomic_save_presentation,
                         write_error, write_json)
from pptx_create import add_bullets
from pptx_text import iter_text_frames, replace_all_text, replace_text


class TemplateError(PptxSkillError):
    code = "invalid_template_spec"


def fill_tokens(prs, values):
    counts = {}
    for token, value in values.items():
        counts[token] = replace_text(prs, "{{%s}}" % token, str(value))
    return counts


def remaining_tokens(prs):
    found = set()
    pattern = re.compile(r"{{([^{}]+)}}")
    for slide in prs.slides:
        for text_frame in iter_text_frames(slide):
            found.update(pattern.findall(text_frame.text))
    return sorted(found)


def find_layout(prs, ref):
    if not isinstance(ref, dict) or set(ref) != {"index", "name"}:
        raise TemplateError(
            "layout must contain exactly index and name",
            code="invalid_layout_ref")
    index = ref["index"]
    if isinstance(index, bool) or not isinstance(index, int):
        raise TemplateError("layout.index must be an integer",
                            code="invalid_layout_ref")
    if index < 0 or index >= len(prs.slide_layouts):
        raise TemplateError(f"layout index is out of range: {index}",
                            code="invalid_layout_ref")
    layout = prs.slide_layouts[index]
    if layout.name != ref["name"]:
        raise TemplateError(
            f"layout name mismatch: expected {ref['name']!r}, "
            f"found {layout.name!r}", code="invalid_layout_ref")
    return layout


def validate_add_spec(spec):
    if not isinstance(spec, dict):
        raise TemplateError("add-slides spec must be an object",
                            code="invalid_type")
    unknown = sorted(set(spec) - {"schema", "slides"})
    if unknown:
        raise TemplateError(f"add-slides spec contains unknown fields: {unknown}",
                            code="unknown_field")
    if spec.get("schema") != "powerpoint.template-add/v1":
        raise TemplateError(
            "add-slides schema must be 'powerpoint.template-add/v1'",
            code="unsupported_schema")
    slides = spec.get("slides")
    if not isinstance(slides, list):
        raise TemplateError("add-slides slides must be an array",
                            code="missing_field")
    for number, slide in enumerate(slides):
        context = f"slides[{number}]"
        if not isinstance(slide, dict):
            raise TemplateError(f"{context} must be an object",
                                code="invalid_type")
        unknown = sorted(set(slide) - {"layout", "title", "bullets", "notes"})
        if unknown:
            raise TemplateError(f"{context} contains unknown fields: {unknown}",
                                code="unknown_field")
        if "layout" not in slide:
            raise TemplateError(f"{context} requires layout",
                                code="missing_field")
        if "bullets" in slide and not isinstance(slide["bullets"], list):
            raise TemplateError(f"{context}.bullets must be an array",
                                code="invalid_type")


def add_slides(prs, spec):
    validate_add_spec(spec)
    added = []
    for slide_spec in spec["slides"]:
        layout = find_layout(prs, slide_spec["layout"])
        slide = prs.slides.add_slide(layout)
        if slide_spec.get("title") is not None and slide.shapes.title:
            replace_all_text(
                slide.shapes.title.text_frame, slide_spec["title"])
        if slide_spec.get("bullets"):
            body = next((
                ph for ph in slide.placeholders
                if ph.placeholder_format.type in {
                    PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT,
                }
            ), None)
            if body is None:
                raise TemplateError(
                    f"layout {layout.name!r} has no body placeholder",
                    code="placeholder_not_found")
            add_bullets(body.text_frame, slide_spec["bullets"])
        if slide_spec.get("notes") is not None:
            replace_all_text(
                slide.notes_slide.notes_text_frame, slide_spec["notes"])
        added.append({"slide_id": int(slide.slide_id),
                      "index": len(prs.slides) - 1})
    return added


def main(argv=None):
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser(
        description="Fill tokens and append slides using strict layout refs.",
        epilog=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("template", help="path to the template .pptx")
    parser.add_argument("output", help="output .pptx path")
    parser.add_argument("--values", metavar="JSON",
                        help="JSON object mapping token to replacement value")
    parser.add_argument("--add-slides", metavar="SPEC_JSON",
                        help="powerpoint.template-add/v1 spec")
    parser.add_argument("--strict", action="store_true",
                        help="fail for unused values or remaining tokens")
    args = parser.parse_args(argv)

    try:
        prs = Presentation(args.template)
        token_counts = {}
        if args.values:
            with open(args.values, encoding="utf-8") as fh:
                values = json.load(fh)
            if not isinstance(values, dict) or any(
                    not isinstance(key, str) or not key for key in values):
                raise TemplateError(
                    "values must be an object with non-empty string keys",
                    code="invalid_type")
            token_counts = fill_tokens(prs, values)
        added = []
        if args.add_slides:
            with open(args.add_slides, encoding="utf-8") as fh:
                added = add_slides(prs, json.load(fh))
        residual = remaining_tokens(prs)
        unused = sorted(token for token, count in token_counts.items() if count == 0)
        if args.strict and unused:
            raise TemplateError(f"unused template values: {unused}",
                                code="unused_token_value")
        if args.strict and residual:
            raise TemplateError(f"unresolved template tokens: {residual}",
                                code="unresolved_token")
        output = atomic_save_presentation(prs, args.output)
    except PptxSkillError as exc:
        write_error(exc, stream=sys.stderr)
        return 2
    except (OSError, KeyError, TypeError, ValueError, AttributeError,
            IndexError, json.JSONDecodeError) as exc:
        write_error(exc, stream=sys.stderr)
        return 2

    write_json({
        "schema": "powerpoint.template-result/v1",
        "ok": True,
        "output": output,
        "tokens_filled": sum(token_counts.values()),
        "token_counts": token_counts,
        "unused_values": unused,
        "remaining_tokens": residual,
        "added_slides": added,
    }, stream=sys.stdout)
    return 0


if __name__ == "__main__":
    sys.exit(main())
