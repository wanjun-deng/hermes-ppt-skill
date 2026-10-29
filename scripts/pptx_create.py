#!/usr/bin/env python3
"""Create a .pptx presentation from a JSON deck spec.

Spec format (all positions/sizes in inches, colors as RRGGBB hex):
{
  "schema": "powerpoint.create/v1",
  "slide_size": "16:9",              // or "4:3" (default "16:9")
  "slides": [
    {"layout": "title", "title": "My Deck", "subtitle": "Q3 review"},
    {"layout": "title_content", "title": "Agenda",
     "bullets": ["Top item",
                 {"text": "Sub item", "level": 1, "bold": true,
                  "size": 18, "color": "CC0000", "font": "Arial",
                  "italic": false,
                  "link": "https://example.com/agenda"}],
     "background": "1F2937",           // solid slide background (hex)
     "footer": "Confidential",         // footer placeholder text
     "slide_number": true,             // enable slide-number placeholder
     "notes": "Speaker notes for this slide"},
    {"layout": "blank", "title": "Widgets",
     "images":  [{"path": "logo.png", "left": 1, "top": 1, "width": 3}],
     "tables":  [{"left": 1, "top": 2, "width": 6, "height": 2,
                  "rows": [["H1", "H2"], ["a", "b"]]}],
     "shapes":  [{"type": "rounded_rectangle", "left": 8, "top": 1,
                  "width": 3, "height": 1, "fill": "4472C4",
                  "text": "Callout", "text_color": "FFFFFF"}],
     "charts":  [{"type": "bar", "left": 1, "top": 3, "width": 6,
                  "height": 3.5, "title": "Sales",
                  "categories": ["Q1", "Q2"],
                  "series": {"North": [10, 20], "South": [7, 13]}}]}
  ]
}
Layouts: title, title_content, section, two_content, title_only, blank
Chart types: bar, bar_h, line, pie   Shape types: rectangle,
rounded_rectangle, oval, diamond, right_arrow, chevron
"""
import argparse
import copy
import json
import os
import re
import sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.util import Inches, Pt

from pptx_common import (PptxSkillError, atomic_save_presentation,
                         write_error, write_json)

LAYOUTS = {"title": 0, "title_content": 1, "section": 2,
           "two_content": 3, "title_only": 5, "blank": 6}
CHART_TYPES = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
               "bar_h": XL_CHART_TYPE.BAR_CLUSTERED,
               "line": XL_CHART_TYPE.LINE_MARKERS,
               "pie": XL_CHART_TYPE.PIE}
SHAPE_TYPES = {"rectangle": MSO_SHAPE.RECTANGLE,
               "rounded_rectangle": MSO_SHAPE.ROUNDED_RECTANGLE,
               "oval": MSO_SHAPE.OVAL, "diamond": MSO_SHAPE.DIAMOND,
               "right_arrow": MSO_SHAPE.RIGHT_ARROW,
               "chevron": MSO_SHAPE.CHEVRON}


class CreateError(PptxSkillError):
    code = "invalid_create_spec"


def _unknown_fields(mapping, allowed, context):
    unknown = sorted(set(mapping) - set(allowed))
    if unknown:
        raise CreateError(f"{context} contains unknown fields: {unknown}",
                          code="unknown_field")


def _number(value, context, *, positive=False):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise CreateError(f"{context} must be numeric", code="invalid_type")
    if positive and value <= 0:
        raise CreateError(f"{context} must be greater than zero",
                          code="invalid_value")


def _color(value, context):
    if not isinstance(value, str) or not re.fullmatch(r"[0-9A-Fa-f]{6}", value):
        raise CreateError(f"{context} must be a six-digit RGB color",
                          code="invalid_color")


def _validate_positioned(item, context, *, require_path=False):
    if require_path:
        if not isinstance(item.get("path"), str) or not item["path"]:
            raise CreateError(f"{context}.path must be a non-empty string",
                              code="missing_field")
        if not os.path.isfile(item["path"]):
            raise CreateError(f"image file not found: {item['path']}",
                              code="input_not_found")
    for field in ("left", "top"):
        if field in item:
            _number(item[field], f"{context}.{field}")
    for field in ("width", "height"):
        if field in item:
            _number(item[field], f"{context}.{field}", positive=True)


def validate_spec(spec):
    if not isinstance(spec, dict):
        raise CreateError("spec root must be an object", code="invalid_type")
    _unknown_fields(spec, {"schema", "slide_size", "slides"}, "spec")
    if spec.get("schema") != "powerpoint.create/v1":
        raise CreateError("spec.schema must be 'powerpoint.create/v1'",
                          code="unsupported_schema")
    if spec.get("slide_size", "16:9") not in {"16:9", "4:3"}:
        raise CreateError("slide_size must be '16:9' or '4:3'",
                          code="invalid_slide_size")
    slides = spec.get("slides")
    if not isinstance(slides, list):
        raise CreateError("spec.slides must be an array", code="missing_field")
    slide_fields = {
        "layout", "title", "subtitle", "bullets", "background", "footer",
        "slide_number", "notes", "images", "tables", "shapes", "charts",
    }
    for slide_number, slide in enumerate(slides):
        context = f"slides[{slide_number}]"
        if not isinstance(slide, dict):
            raise CreateError(f"{context} must be an object", code="invalid_type")
        _unknown_fields(slide, slide_fields, context)
        layout = slide.get("layout", "title_content")
        if layout not in LAYOUTS:
            raise CreateError(f"{context}.layout is unknown: {layout!r}",
                              code="invalid_layout_ref")
        for field in ("title", "subtitle", "footer", "notes"):
            if (field in slide and slide[field] is not None
                    and not isinstance(slide[field], str)):
                raise CreateError(
                    f"{context}.{field} must be a string or null",
                    code="invalid_type")
        if "background" in slide:
            _color(slide["background"], f"{context}.background")
        if "slide_number" in slide and not isinstance(slide["slide_number"], bool):
            raise CreateError(f"{context}.slide_number must be boolean",
                              code="invalid_type")
        bullets = slide.get("bullets", [])
        if not isinstance(bullets, list):
            raise CreateError(f"{context}.bullets must be an array",
                              code="invalid_type")
        for item_number, item in enumerate(bullets):
            item_context = f"{context}.bullets[{item_number}]"
            if isinstance(item, str):
                continue
            if not isinstance(item, dict):
                raise CreateError(f"{item_context} must be a string or object",
                                  code="invalid_type")
            _unknown_fields(item, {"text", "level", "size", "bold", "italic",
                                   "font", "color", "link"}, item_context)
            if "text" not in item or not isinstance(item["text"], str):
                raise CreateError(f"{item_context}.text must be a string",
                                  code="missing_field")
            for field in ("bold", "italic"):
                if field in item and not isinstance(item[field], bool):
                    raise CreateError(f"{item_context}.{field} must be boolean",
                                      code="invalid_type")
            for field in ("font", "link"):
                if field in item and not isinstance(item[field], str):
                    raise CreateError(f"{item_context}.{field} must be a string",
                                      code="invalid_type")
            if "level" in item and (isinstance(item["level"], bool)
                                    or not isinstance(item["level"], int)
                                    or item["level"] not in range(5)):
                raise CreateError(f"{item_context}.level must be 0..4",
                                  code="invalid_value")
            if "size" in item:
                _number(item["size"], f"{item_context}.size", positive=True)
            if "color" in item:
                _color(item["color"], f"{item_context}.color")
        images = slide.get("images", [])
        if not isinstance(images, list):
            raise CreateError(f"{context}.images must be an array",
                              code="invalid_type")
        for number, item in enumerate(images):
            item_context = f"{context}.images[{number}]"
            if not isinstance(item, dict):
                raise CreateError(f"{item_context} must be an object",
                                  code="invalid_type")
            _unknown_fields(item, {"path", "left", "top", "width", "height"},
                            item_context)
            _validate_positioned(item, item_context, require_path=True)
        tables = slide.get("tables", [])
        if not isinstance(tables, list):
            raise CreateError(f"{context}.tables must be an array",
                              code="invalid_type")
        for number, item in enumerate(tables):
            item_context = f"{context}.tables[{number}]"
            if not isinstance(item, dict):
                raise CreateError(f"{item_context} must be an object",
                                  code="invalid_type")
            _unknown_fields(item, {"rows", "left", "top", "width", "height"},
                            item_context)
            _validate_positioned(item, item_context)
            rows = item.get("rows")
            if (not isinstance(rows, list) or not rows
                    or not isinstance(rows[0], list) or not rows[0]):
                raise CreateError(f"{item_context}.rows must be a non-empty grid",
                                  code="invalid_table")
            width = len(rows[0])
            if any(not isinstance(row, list) or len(row) != width for row in rows):
                raise CreateError(f"{item_context}.rows must be rectangular",
                                  code="invalid_table")
        shapes = slide.get("shapes", [])
        if not isinstance(shapes, list):
            raise CreateError(f"{context}.shapes must be an array",
                              code="invalid_type")
        for number, item in enumerate(shapes):
            item_context = f"{context}.shapes[{number}]"
            if not isinstance(item, dict):
                raise CreateError(f"{item_context} must be an object",
                                  code="invalid_type")
            _unknown_fields(item, {"type", "left", "top", "width", "height",
                                   "fill", "text", "text_color"}, item_context)
            kind = item.get("type", "rectangle")
            if kind not in SHAPE_TYPES:
                raise CreateError(f"{item_context}.type is unknown: {kind!r}",
                                  code="invalid_shape_type")
            if "text" in item and not isinstance(item["text"], str):
                raise CreateError(
                    f"{item_context}.text must be a string",
                    code="invalid_type")
            _validate_positioned(item, item_context)
            for field in ("fill", "text_color"):
                if field in item:
                    _color(item[field], f"{item_context}.{field}")
        charts = slide.get("charts", [])
        if not isinstance(charts, list):
            raise CreateError(f"{context}.charts must be an array",
                              code="invalid_type")
        for number, item in enumerate(charts):
            item_context = f"{context}.charts[{number}]"
            if not isinstance(item, dict):
                raise CreateError(f"{item_context} must be an object",
                                  code="invalid_type")
            _unknown_fields(item, {"type", "left", "top", "width", "height",
                                   "title", "categories", "series"}, item_context)
            kind = item.get("type", "bar")
            if kind not in CHART_TYPES:
                raise CreateError(f"{item_context}.type is unknown: {kind!r}",
                                  code="invalid_chart_type")
            _validate_positioned(item, item_context)
            categories = item.get("categories")
            series = item.get("series")
            if not isinstance(categories, list) or not categories:
                raise CreateError(f"{item_context}.categories must be non-empty",
                                  code="invalid_chart")
            if not isinstance(series, dict) or not series:
                raise CreateError(f"{item_context}.series must be non-empty",
                                  code="invalid_chart")
            if any(not isinstance(name, str) or not name for name in series):
                raise CreateError(
                    f"{item_context}.series names must be non-empty strings",
                    code="invalid_chart")
            if any(not isinstance(values, list)
                   or len(values) != len(categories) for values in series.values()):
                raise CreateError(f"{item_context}.series must match categories",
                                  code="invalid_chart")


def style_run(run, spec):
    """Apply font styling from a bullet/text spec dict to a run."""
    font = run.font
    if spec.get("size"):
        font.size = Pt(spec["size"])
    if spec.get("bold") is not None:
        font.bold = spec["bold"]
    if spec.get("italic") is not None:
        font.italic = spec["italic"]
    if spec.get("font"):
        font.name = spec["font"]
    if spec.get("color"):
        font.color.rgb = RGBColor.from_string(spec["color"])
    if spec.get("link"):
        run.hyperlink.address = spec["link"]


def add_bullets(text_frame, bullets):
    text_frame.clear()
    for i, item in enumerate(bullets):
        if isinstance(item, str):
            item = {"text": item}
        para = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
        para.level = int(item.get("level", 0))
        run = para.add_run()
        run.text = item.get("text", "")
        style_run(run, item)


def copy_layout_placeholder(slide, ph_idx):
    """Copy a layout placeholder (footer=11, slide number=12) onto the
    slide so it actually renders; returns the shape or None if the layout
    does not provide it."""
    for ph in slide.slide_layout.placeholders:
        if ph.placeholder_format.idx == ph_idx:
            slide.shapes._spTree.append(copy.deepcopy(ph._element))
            for shape in slide.placeholders:
                if shape.placeholder_format.idx == ph_idx:
                    return shape
    return None


def build_slide(prs, spec):
    layout_idx = LAYOUTS[spec.get("layout", "title_content")]
    slide = prs.slides.add_slide(prs.slide_layouts[layout_idx])

    if spec.get("background"):
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor.from_string(spec["background"])
    if spec.get("slide_number"):
        if copy_layout_placeholder(slide, 12) is None:
            raise CreateError(
                f"layout {slide.slide_layout.name!r} has no slide-number "
                "placeholder", code="placeholder_not_found")
    if spec.get("footer") is not None:
        shape = copy_layout_placeholder(slide, 11)
        if shape is None:
            raise CreateError(
                f"layout {slide.slide_layout.name!r} has no footer "
                "placeholder", code="placeholder_not_found")
        shape.text_frame.text = spec["footer"]

    if spec.get("title") is not None:
        title_shape = slide.shapes.title
        if title_shape is None and spec.get("layout") == "blank":
            title_shape = slide.shapes.add_textbox(
                Inches(0.5), Inches(0.25),
                prs.slide_width - Inches(1), Inches(0.75),
            )
        if title_shape is None:
            raise CreateError(
                f"layout {slide.slide_layout.name!r} has no title placeholder",
                code="placeholder_not_found")
        title_shape.text = spec["title"]
    if spec.get("subtitle") is not None:
        subtitle = next((
            ph for ph in slide.placeholders
            if ph.placeholder_format.type == PP_PLACEHOLDER.SUBTITLE
        ), None)
        if subtitle is None:
            raise CreateError(
                f"layout {slide.slide_layout.name!r} has no subtitle "
                "placeholder", code="placeholder_not_found")
        subtitle.text = spec["subtitle"]
    if spec.get("bullets"):
        body = next((
            ph for ph in slide.placeholders
            if ph.placeholder_format.type in {
                PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT,
            }
        ), None)
        if body is None:
            body = slide.shapes.add_textbox(Inches(0.5), Inches(1.5),
                                            Inches(9), Inches(5))
        add_bullets(body.text_frame, spec["bullets"])

    for img in spec.get("images", []):
        kwargs = {}
        if img.get("width"):
            kwargs["width"] = Inches(img["width"])
        if img.get("height"):
            kwargs["height"] = Inches(img["height"])
        slide.shapes.add_picture(img["path"], Inches(img.get("left", 1)),
                                 Inches(img.get("top", 1)), **kwargs)

    for tbl in spec.get("tables", []):
        rows = tbl["rows"]
        shape = slide.shapes.add_table(
            len(rows), len(rows[0]), Inches(tbl.get("left", 1)),
            Inches(tbl.get("top", 2)), Inches(tbl.get("width", 6)),
            Inches(tbl.get("height", 2)))
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                shape.table.cell(r, c).text = str(val)

    for shp in spec.get("shapes", []):
        shape = slide.shapes.add_shape(
            SHAPE_TYPES[shp.get("type", "rectangle")],
            Inches(shp.get("left", 1)), Inches(shp.get("top", 1)),
            Inches(shp.get("width", 2)), Inches(shp.get("height", 1)))
        if shp.get("fill"):
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(shp["fill"])
        if "text" in shp:
            shape.text_frame.text = shp["text"]
            if shp.get("text_color"):
                run = shape.text_frame.paragraphs[0].runs[0]
                run.font.color.rgb = RGBColor.from_string(shp["text_color"])

    for cht in spec.get("charts", []):
        data = CategoryChartData()
        data.categories = cht["categories"]
        for name, values in cht["series"].items():
            data.add_series(name, values)
        frame = slide.shapes.add_chart(
            CHART_TYPES[cht.get("type", "bar")],
            Inches(cht.get("left", 1)), Inches(cht.get("top", 2)),
            Inches(cht.get("width", 6)), Inches(cht.get("height", 4)), data)
        if cht.get("title"):
            frame.chart.has_title = True
            frame.chart.chart_title.text_frame.text = cht["title"]

    if spec.get("notes") is not None:
        slide.notes_slide.notes_text_frame.text = spec["notes"]
    return slide


def main(argv=None):
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser(
        description="Create a .pptx deck from a powerpoint.create/v1 spec.",
        epilog=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("spec", help="path to JSON deck spec")
    parser.add_argument("output", help="output .pptx path")
    args = parser.parse_args(argv)

    try:
        with open(args.spec, encoding="utf-8") as fh:
            spec = json.load(fh)
        validate_spec(spec)
        prs = Presentation()
        if spec.get("slide_size", "16:9") == "16:9":
            prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        else:
            prs.slide_width, prs.slide_height = Inches(10), Inches(7.5)
        for slide_spec in spec["slides"]:
            build_slide(prs, slide_spec)
        output = atomic_save_presentation(prs, args.output)
    except PptxSkillError as exc:
        write_error(exc, stream=sys.stderr)
        return 2
    except (OSError, KeyError, TypeError, ValueError, AttributeError,
            IndexError, json.JSONDecodeError) as exc:
        write_error(exc, stream=sys.stderr)
        return 2

    write_json({
        "schema": "powerpoint.create-result/v1",
        "ok": True,
        "output": output,
        "slides": len(prs.slides),
    }, stream=sys.stdout)
    return 0


if __name__ == "__main__":
    sys.exit(main())
