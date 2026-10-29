#!/usr/bin/env python3
"""Advisory geometry checks; rendered pages remain the authority on text fit."""
import argparse
import json
from pathlib import Path
import sys

from pptx import Presentation
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from pptx_common import write_error
from pptx_inventory import recommended_ref, sha256_file, walk_shapes

TOLERANCE = 18288  # 0.02 inch; report, do not silently move or resize anything.


def _opaque_rectangle(shape):
    if shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE or shape.auto_shape_type != MSO_SHAPE.RECTANGLE:
        return False
    if shape.fill.type != MSO_FILL_TYPE.SOLID:
        return False
    fill = shape._element.spPr.find(qn("a:solidFill"))
    if fill is None:
        return False
    if any(e.tag in {qn("a:alphaMod"), qn("a:alphaOff")} for e in fill.iter()):
        return False
    return all(int(e.get("val", "0")) == 100000 for e in fill.iter(qn("a:alpha")))


def _coverage(first, second):
    x, y, w, h = first
    a, b, c, d = second
    intersection = max(0, min(x+w, a+c)-max(x, a)) * max(0, min(y+h, b+d)-max(y, b))
    return intersection / (w*h) if w > 0 and h > 0 else 0


def _inspect(path):
    prs = Presentation(path)
    findings, slide_ids = [], []
    for index, slide in enumerate(prs.slides):
        slide_ids.append(slide.slide_id)
        nodes = [n for n in walk_shapes(slide.shapes) if n.shape.shape_type != MSO_SHAPE_TYPE.GROUP]
        def issue(node, code, evidence, other=None):
            row = {"slide_index": index, "shape": recommended_ref(node), "code": code,
                   "severity": "warning", "evidence": evidence}
            if other is not None:
                row["related_shape"] = recommended_ref(other)
            findings.append(row)
        for i, node in enumerate(nodes):
            shape = node.shape
            bounds = list(node.slide_bounds)
            if node.bounds_accuracy != "exact":
                issue(node, "geometry_unconfirmed", {"bounds_emu": bounds,
                      "reason": "rotation/flip bounds are approximate; inspect the render"})
                continue
            x, y, w, h = bounds
            if x < -TOLERANCE or y < -TOLERANCE or x+w > prs.slide_width+TOLERANCE or y+h > prs.slide_height+TOLERANCE:
                issue(node, "off_canvas", {"bounds_emu": bounds,
                      "slide_size_emu": [int(prs.slide_width), int(prs.slide_height)],
                      "reason": "may be intentional decorative bleed; verify visible content"})
            meaningful = ((shape.has_text_frame and shape.text.strip()) or shape.has_table
                          or shape.has_chart or shape._element.tag == qn("p:pic"))
            if not meaningful:
                continue
            for other in nodes[i+1:]:
                if other.bounds_accuracy != "exact" or not _opaque_rectangle(other.shape):
                    continue
                ratio = _coverage(bounds, other.slide_bounds)
                if ratio >= 0.95:
                    issue(node, "likely_occluded", {"bounds_emu": bounds,
                          "occluder_bounds_emu": list(other.slide_bounds),
                          "coverage": round(ratio, 4),
                          "reason": "later solid rectangle covers at least 95% of the object"}, other)
                    break
    return findings, slide_ids


def _signature(row):
    return json.dumps({k: row[k] for k in ("code", "shape", "related_shape", "evidence") if k in row},
                      sort_keys=True, ensure_ascii=False)


def geometry_report(path, baseline=None, source_slide_indexes=None):
    path = Path(path).resolve()
    rows, ids = _inspect(path)
    old, old_ids = _inspect(baseline) if baseline else ([], [])
    if source_slide_indexes is not None and (
        len(source_slide_indexes) != len(ids) or any(type(i) is not int or i < 0 or i >= len(old_ids) for i in source_slide_indexes)
    ):
        raise ValueError("geometry baseline slide mapping is invalid")
    by_id = {sid: i for i, sid in enumerate(old_ids)}
    previous = {(r["slide_index"], _signature(r)) for r in old}
    for row in rows:
        index = row["slide_index"]
        source_index = source_slide_indexes[index] if source_slide_indexes is not None else by_id.get(ids[index])
        row["origin"] = "baseline" if (source_index, _signature(row)) in previous else "new"
    return {"schema": "powerpoint.geometry/v1", "artifact_sha256": sha256_file(path),
            "baseline_sha256": sha256_file(baseline) if baseline else None,
            "policy": "advisory", "tolerance_emu": TOLERANCE, "findings": rows,
            "new_finding_count": sum(r["origin"] == "new" for r in rows),
            "limitations": ["does not prove text fit or visual correctness",
                           "checks slide objects; inherited master/layout artwork requires visual review",
                           "occlusion checks only explicit opaque rectangles",
                           "approximate bounds are reported for visual review, not treated as exact defects"]}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx")
    parser.add_argument("--baseline")
    parser.add_argument("--output")
    parser.add_argument("--fail-on-new", action="store_true", help="explicit opt-in to fail on advisory findings")
    args = parser.parse_args(argv)
    try:
        result = geometry_report(args.pptx, args.baseline)
        if args.output:
            from pptx_workflow import save_json
            save_json(args.output, result)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 1 if args.fail_on_new and result["new_finding_count"] else 0
    except Exception as exc:
        write_error(exc, stream=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
