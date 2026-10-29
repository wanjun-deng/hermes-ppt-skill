"""Optional intent plans must affect content acceptance without restyling."""
import json
from pathlib import Path
import sys
import pytest
from pptx import Presentation
from pptx.util import Inches

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))
from pptx_design import review_design
from pptx_inventory import inspect_presentation, sha256_file
from pptx_workflow import build


def test_design_requires_matching_unique_page_ids():
    design = {"schema": "powerpoint.design/v1", "objective": "Explain progress",
              "branding": "template", "pages": [
                  {"id": "wrong", "purpose": "Summarize", "focus": "Progress", "kind": "content"}]}
    with pytest.raises(ValueError, match="IDs"):
        review_design(design, [{"id": "page1"}])
    design["pages"] = [dict(design["pages"][0], id="page1")] * 2
    with pytest.raises(ValueError, match="IDs"):
        review_design(design, [{"id": "page1"}])


def test_build_keeps_design_requirements_and_archives_revision(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text = "Sample"
    template = tmp_path / "template.pptx"
    prs.save(template)
    inv = inspect_presentation(template)
    spec = {"schema": "powerpoint.presentation/v1",
            "template": {"path": "template.pptx", "sha256": sha256_file(template), "slide_count": 1},
            "pages": [{"id": "page1", "source_slide": inv["slides"][0]["recommended_ref"],
                       "text": [{"shape": inv["slides"][0]["shapes"][0]["recommended_ref"], "value": "Actual progress"}]}]}
    plan = tmp_path / "plan.json"
    plan.write_text(json.dumps(spec), encoding="utf-8")
    design = {"schema": "powerpoint.design/v1", "objective": "Explain progress", "branding": "template",
              "colors": ["123456"], "pages": [
                  {"id": "page1", "purpose": "Summarize", "focus": "Progress", "kind": "content",
                   "density": "high", "must_keep": ["Missing required fact"]}]}
    design_path = tmp_path / "design.json"
    design_path.write_text(json.dumps(design), encoding="utf-8")
    folder = tmp_path / "run"
    build(plan, folder, False, design_path)
    report = json.loads((folder / "qa-report.json").read_text())
    check = next(row for row in report["content_checks"] if row["id"] == "design:page1")
    assert not check["passed"] and check["missing"] == ["Missing required fact"]
    compilation = json.loads((folder / "compilation.json").read_text())
    assert any(row["code"] == "design_density_review" for row in compilation["warnings"])
    geometry = json.loads((folder / "geometry-report.json").read_text())
    assert geometry["artifact_sha256"] == sha256_file(folder / "output.pptx")
    assert geometry["baseline_sha256"] == sha256_file(template)
    old_hash = sha256_file(folder / "output.pptx")
    # Sidecar colors record intent; they do not change the template's fill.
    assert Presentation(folder / "output.pptx").slides[0].background.fill.type == Presentation(template).slides[0].background.fill.type
    build(plan, folder, False)
    revision = folder / "revisions" / old_hash
    assert (revision / "design-plan.json").is_file()
    assert (revision / "geometry-report.json").is_file()
    assert (revision / "presentation-plan.json").is_file()
    assert not (folder / "design-plan.json").exists()
    assert Path(json.loads((folder / "presentation-plan.json").read_text())["template"]["path"]).is_absolute()
