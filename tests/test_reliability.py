"""Behavioral regressions for scope proof, geometry, fonts, and mixed runs."""
import json
from pathlib import Path
import random
import subprocess
import sys
from types import SimpleNamespace

from jsonschema import Draft202012Validator
import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))
from pptx_constrained_edit import execute_edit
from pptx_fonts import probe_font, set_run_fonts, set_presentation_fonts
from pptx_geometry import geometry_report
from pptx_inventory import inspect_presentation, sha256_file
from pptx_text import replace_in_text_frame


def validate(name, result):
    schema = json.loads((SCRIPTS.parent / "schemas" / name).read_text(encoding="utf-8"))
    Draft202012Validator(schema).validate(result)


@pytest.fixture
def source(tmp_path):
    prs = Presentation()
    for index in range(2):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        p = box.text_frame.paragraphs[0]
        for text, bold in (("Revenue: 1", True), ("50 万元 unchanged", False)):
            run = p.add_run()
            run.text = text
            run.font.bold = bold
            run.font.size = Pt(18)
        p.runs[1].hyperlink.address = "https://example.com"
    table = prs.slides[1].shapes.add_table(1, 2, Inches(1), Inches(3), Inches(5), Inches(1)).table
    table.cell(0, 0).text = "Q1"
    table.cell(0, 1).text = "120"
    path = tmp_path / "source.pptx"
    prs.save(path)
    return path


def contract(source, mode, action, **extra):
    inv = inspect_presentation(source)
    op = {"action": action, "slide": inv["slides"][0]["recommended_ref"],
          "shape": inv["slides"][0]["shapes"][0]["recommended_ref"], **extra}
    if action == "add_text":
        op.pop("shape")
    if action == "append_slide":
        op.pop("shape")
        op.pop("slide")
    return {"schema": "powerpoint.edit/v1",
            "source": {"sha256": sha256_file(source), "slide_count": 2},
            "mode": mode, "operations": [op]}


def test_mixed_run_replacement_preserves_suffix_and_hyperlink(source, tmp_path):
    spec = contract(source, "precise", "replace_text", old="150", new="180")
    output = tmp_path / "out.pptx"
    result = execute_edit(source, spec, output)
    validate("powerpoint-edit-proof-v1.schema.json", result)
    validate("powerpoint-edit-v1.schema.json", spec)
    p = Presentation(output).slides[0].shapes[0].text_frame.paragraphs[0]
    assert p.text == "Revenue: 180 万元 unchanged"
    assert p.runs[0].font.bold is True
    assert p.runs[1].font.bold is False
    assert p.runs[1].hyperlink.address == "https://example.com"
    assert result["package_changes"]["changed"] == ["ppt/slides/slide1.xml"]
    assert result["scope_verified"] and result["visual_review_pending"]
    assert sha256_file(source) == spec["source"]["sha256"]


def test_replacement_matches_python_across_arbitrary_run_boundaries():
    rng = random.Random(41)
    for _ in range(100):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        frame = slide.shapes.add_textbox(0, 0, Inches(5), Inches(1)).text_frame
        text = "aa150bb150150cc"
        cursor = 0
        while cursor < len(text):
            size = rng.randint(1, 4)
            run = frame.paragraphs[0].add_run()
            run.text = text[cursor:cursor+size]
            run.font.bold = bool(cursor % 2)
            cursor += size
        assert replace_in_text_frame(frame, "150", "8") == 3
        assert frame.text == text.replace("150", "8")


def test_replacement_does_not_cross_line_break():
    prs = Presentation()
    frame = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(0, 0, Inches(5), Inches(1)).text_frame
    p = frame.paragraphs[0]
    p.add_run().text = "15"
    p.add_line_break()
    p.add_run().text = "0"
    assert replace_in_text_frame(frame, "150", "180") == 0
    assert frame.text == "15\v0"


@pytest.mark.parametrize("mode,action,values", [
    ("restyle", "replace_text", {"old": "150", "new": "180"}),
    ("enrich", "set_font", {"east_asian": "Noto Sans SC"}),
    ("precise", "add_text", {"text": "extra", "bounds": {"left": 1, "top": 4, "width": 3, "height": 1}}),
])
def test_modes_reject_forbidden_operations(source, tmp_path, mode, action, values):
    with pytest.raises(ValueError, match="forbidden"):
        execute_edit(source, contract(source, mode, action, **values), tmp_path / "out.pptx")
    assert not (tmp_path / "out.pptx").exists()


def test_invalid_count_hash_ref_and_unknown_fields_do_not_replace_output(source, tmp_path):
    output = tmp_path / "out.pptx"
    output.write_bytes(source.read_bytes())
    original = sha256_file(output)
    spec = contract(source, "precise", "replace_text", old="150", new="180", expected_count=2)
    with pytest.raises(ValueError, match="count"):
        execute_edit(source, spec, output)
    spec["operations"][0]["expected_count"] = 1
    spec["source"]["sha256"] = "0"*64
    with pytest.raises(ValueError, match="hash"):
        execute_edit(source, spec, output)
    spec["source"]["sha256"] = sha256_file(source)
    spec["operations"][0]["slide"]["slide_id"] += 1
    with pytest.raises(ValueError, match="locator"):
        execute_edit(source, spec, output)
    spec["operations"][0]["unknown"] = True
    with pytest.raises(Exception, match="schema|properties|valid"):
        execute_edit(source, spec, output)
    assert sha256_file(output) == original


def test_font_only_edit_keeps_words_geometry_and_other_properties(source, tmp_path):
    spec = contract(source, "restyle", "set_font", east_asian="WenQuanYi Zen Hei")
    output = tmp_path / "font.pptx"
    proof = execute_edit(source, spec, output)
    before, after = Presentation(source).slides[0].shapes[0], Presentation(output).slides[0].shapes[0]
    assert before.text == after.text
    assert (before.left, before.top, before.width, before.height) == (after.left, after.top, after.width, after.height)
    for p in after.text_frame.paragraphs:
        for run in p.runs:
            assert run._r.rPr.find(qn("a:ea")).get("typeface") == "WenQuanYi Zen Hei"
    assert proof["package_changes"]["changed"] == ["ppt/slides/slide1.xml"]


def test_frozen_target_property_is_checked(source, tmp_path, monkeypatch):
    import pptx_constrained_edit as editor
    original = editor.set_shape_fonts
    def broken(shape, **fonts):
        original(shape, **fonts)
        shape.left += Inches(1)
    monkeypatch.setattr(editor, "set_shape_fonts", broken)
    with pytest.raises(ValueError, match="frozen"):
        execute_edit(source, contract(source, "restyle", "set_font", east_asian="Test"), tmp_path / "out.pptx")
    assert not (tmp_path / "out.pptx").exists()


def test_package_changes_are_validated_before_atomic_publication(source, tmp_path, monkeypatch):
    import pptx_constrained_edit as editor
    original = editor.atomic_save_presentation
    def broken(prs, output, *, validator):
        prs.core_properties.title = "unauthorized metadata edit"
        return original(prs, output, validator=validator)
    monkeypatch.setattr(editor, "atomic_save_presentation", broken)
    output = tmp_path / "out.pptx"
    output.write_bytes(source.read_bytes())
    saved_hash = sha256_file(output)
    with pytest.raises(ValueError, match="out-of-scope"):
        execute_edit(source, contract(source, "precise", "replace_text", old="150", new="180"), output)
    assert sha256_file(output) == saved_hash


@pytest.mark.parametrize("action,extra", [
    ("set_text_style", {"font_size_pt": 22, "bold": False, "color": "123456", "alignment": "center"}),
    ("set_geometry", {"bounds": {"left": 2, "top": 2, "width": 4, "height": 1}}),
])
def test_restyle_preserves_content(source, tmp_path, action, extra):
    output = tmp_path / "styled.pptx"
    execute_edit(source, contract(source, "restyle", action, **extra), output)
    assert Presentation(output).slides[0].shapes[0].text == Presentation(source).slides[0].shapes[0].text


def test_enrich_adds_without_changing_original_objects(source, tmp_path):
    output = tmp_path / "enriched.pptx"
    spec = contract(source, "enrich", "add_text", text="New explanation",
                    bounds={"left": 1, "top": 4, "width": 4, "height": 1})
    proof = execute_edit(source, spec, output)
    assert len(Presentation(output).slides[0].shapes) == len(Presentation(source).slides[0].shapes)+1
    assert proof["package_changes"]["changed"] == ["ppt/slides/slide1.xml"]


def test_enrich_appends_slide_preserving_original_order(source, tmp_path):
    prs = Presentation(source)
    spec = contract(source, "enrich", "append_slide", layout={"index": 6, "name": prs.slide_layouts[6].name},
                    text=[{"text": "Action plan", "bounds": {"left": 1, "top": 1, "width": 5, "height": 1}}])
    proof = execute_edit(source, spec, tmp_path / "append.pptx")
    result = Presentation(proof["output"])
    assert len(result.slides) == 3
    assert [s.slide_id for s in result.slides][:2] == [s.slide_id for s in prs.slides]
    assert not proof["package_changes"]["removed"]


def test_geometry_nonstandard_size_z_order_and_baseline(tmp_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # A background under text must not be called an occluder.
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(10), Inches(7.5)).fill.solid()
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "Visible"
    bleed = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(-1), 0, Inches(2), Inches(1))
    bleed.fill.solid()
    source = tmp_path / "before.pptx"
    prs.save(source)
    assert not any(r["code"] == "likely_occluded" for r in geometry_report(source)["findings"])
    # Later opaque rectangle hides text.
    cover = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(4), Inches(1))
    cover.fill.solid()
    rotated = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(12), Inches(1), Inches(2), Inches(1))
    rotated.rotation = 45
    output = tmp_path / "after.pptx"
    prs.save(output)
    report = geometry_report(output, source)
    validate("powerpoint-geometry-v1.schema.json", report)
    assert any(r["code"] == "off_canvas" and r["origin"] == "baseline" for r in report["findings"])
    assert any(r["code"] == "likely_occluded" and r["origin"] == "new" for r in report["findings"])
    assert any(r["code"] == "geometry_unconfirmed" for r in report["findings"])
    assert not any(r["code"] == "off_canvas" and r["shape"].get("shape_id") == rotated.shape_id for r in report["findings"])


def test_transparent_overlay_is_not_opaque(tmp_path):
    from pptx.oxml.xmlchemy import OxmlElement
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text = "Visible"
    cover = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(4), Inches(1))
    cover.fill.solid()
    cover.fill.fore_color.rgb = RGBColor(0, 0, 0)
    alpha = OxmlElement("a:alpha")
    alpha.set("val", "50000")
    cover._element.spPr.find(qn("a:solidFill"))[0].append(alpha)
    path = tmp_path / "transparent.pptx"
    prs.save(path)
    assert not any(r["code"] == "likely_occluded" for r in geometry_report(path)["findings"])


def test_font_probe_records_substitution_and_actual_glyph_coverage(monkeypatch):
    import pptx_fonts
    monkeypatch.setattr(pptx_fonts.shutil, "which", lambda _: "/usr/bin/fc-match")
    monkeypatch.setattr(pptx_fonts.subprocess, "run", lambda *a, **k: SimpleNamespace(stdout="Actual Font\n/fonts/font.ttf\n20-7e\n"))
    result = probe_font("Missing Font", "A中")
    assert result["status"] == "substituted"
    assert result["missing_characters"] == ["中"] and not result["glyphs_available"]
    validate("powerpoint-font-probe-v1.schema.json", {"schema": "powerpoint.font-probe/v1", "fonts": [result]})
    monkeypatch.setattr(pptx_fonts.shutil, "which", lambda _: None)
    assert probe_font("Anything")["status"] == "unconfirmed"


def test_font_slots_are_ordered_and_creation_covers_table_chart():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(0, 0, Inches(4), Inches(1))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "中文"
    run.hyperlink.address = "https://example.com"
    set_run_fonts(run, latin="Arial", east_asian="WenQuanYi Zen Hei", complex_script="Arial")
    tags = [etree_tag.rsplit("}", 1)[-1] for etree_tag in [e.tag for e in run._r.rPr]]
    assert tags.index("latin") < tags.index("ea") < tags.index("cs") < tags.index("hlinkClick")
    table = slide.shapes.add_table(1, 1, 0, Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "表格"
    data = CategoryChartData()
    data.categories = ["季度"]
    data.add_series("营收", [120])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 0, Inches(3), Inches(4), Inches(3), data).chart
    set_presentation_fonts(prs, east_asian="WenQuanYi Zen Hei")
    assert table.cell(0, 0).text_frame.paragraphs[0].runs[0]._r.rPr.find(qn("a:ea")) is not None
    assert chart.font._rPr.find(qn("a:ea")).get("typeface") == "WenQuanYi Zen Hei"


def test_constrained_cli_emits_auditable_proof(source, tmp_path):
    spec = tmp_path / "edit.json"
    spec.write_text(json.dumps(contract(source, "precise", "replace_text", old="150", new="180")), encoding="utf-8")
    output, report = tmp_path / "output.pptx", tmp_path / "proof.json"
    process = subprocess.run([sys.executable, str(SCRIPTS / "pptx_constrained_edit.py"), str(source),
                              str(spec), str(output), "--report", str(report)], capture_output=True, text=True)
    assert process.returncode == 0, process.stderr
    proof = json.loads(report.read_text(encoding="utf-8"))
    assert proof["artifact_sha256"] == sha256_file(output)
    assert proof["scope_verified"] and proof["visual_review_pending"]

def test_precise_table_geometry_preserves_native_data(source, tmp_path):
    inv = inspect_presentation(source)
    spec = contract(source, "precise", "set_geometry",
                    bounds={"left": 2, "top": 3, "width": 5, "height": 1})
    spec["operations"][0].update(
        slide=inv["slides"][1]["recommended_ref"],
        shape=next(s["recommended_ref"] for s in inv["slides"][1]["shapes"] if s["capabilities"]["table"]))
    output = tmp_path / "table-position.pptx"
    execute_edit(source, spec, output)
    table = Presentation(output).slides[1].shapes[1].table
    assert [[c.text for c in row.cells] for row in table.rows] == [["Q1", "120"]]


def test_geometry_applies_group_scaling_to_leaf_bounds(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(1)).text = "Grouped label"
    group.left, group.width = Inches(9), Inches(3)
    path = tmp_path / "group.pptx"
    prs.save(path)
    findings = geometry_report(path)["findings"]
    row = next(r for r in findings if r["code"] == "off_canvas")
    assert row["shape"]["path"] == [0, 0]
    assert row["evidence"]["bounds_emu"][0] == Inches(9)
    assert row["evidence"]["bounds_emu"][2] == Inches(3)
