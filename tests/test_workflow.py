import importlib.util
import json
from pathlib import Path
import sys

import pytest
from pptx import Presentation
from pptx.util import Inches

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))
from configure_hermes import configure
from pptx_inventory import inspect_presentation, sha256_file
from pptx_materials import extract_materials
from pptx_quality import content_checks, quality_report, VISUAL_CHECKS
from pptx_templates import select_templates
from pptx_workflow import compile_presentation, build


@pytest.fixture
def template(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = "Original heading"
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1)).text = "OLD SAMPLE"
    slide.notes_slide.notes_text_frame.text = "NOTES ONLY"
    path = tmp_path / "template.pptx"
    prs.save(path)
    return path


def spec_for(path):
    inv = inspect_presentation(str(path))
    return {"schema": "powerpoint.presentation/v1", "template": {"path": str(path), "sha256": inv["template"]["sha256"], "slide_count": 1},
            "pages": [{"id": "page1", "source_slide": inv["slides"][0]["recommended_ref"], "text": [{"shape": inv["slides"][0]["shapes"][0]["recommended_ref"], "value": "Real heading"}]}],
            "requirements": [{"id": "heading", "text": ["Real heading"], "slide_indexes": [0]}], "forbidden_text": ["OLD SAMPLE"]}


def test_configuration_migration_is_narrow_and_idempotent(tmp_path):
    (tmp_path / "config.yaml").write_text("model: {default: preserved}\nskills: {disabled: [powerpoint]}\nother: 42\n")
    folder = tmp_path / "skills/productivity"
    folder.mkdir(parents=True)
    (folder / "powerpoint-dev").symlink_to(tmp_path / "missing")
    first = configure(tmp_path, SCRIPTS.parents[1], True)
    second = configure(tmp_path, SCRIPTS.parents[1], True)
    import yaml
    config = yaml.safe_load((tmp_path / "config.yaml").read_text())
    assert config["model"] == {"default": "preserved"} and config["other"] == 42
    assert "pptx-template-generation" in config["skills"]["disabled"]
    assert Path(first["backup"], "powerpoint-dev.broken-link").is_symlink()
    assert second["backup"] is None and not second["config_changed"]


def test_notes_do_not_satisfy_visible_requirement(template):
    checks = content_checks(inspect_presentation(str(template)), [{"id": "visible", "text": ["NOTES ONLY"]}])
    assert not checks[0]["passed"]


def test_evaluator_uses_actual_slide_layouts(template):
    evaluator = SCRIPTS.parents[1] / "powerpoint-skill-eval/scripts/evaluate_run.py"
    module_spec = importlib.util.spec_from_file_location("local_evaluator", evaluator)
    module = importlib.util.module_from_spec(module_spec)
    module_spec.loader.exec_module(module)
    checks = []
    case = {"expected": {"required_layout_names": ["Title Slide"], "required_text": ["NOTES ONLY"]}}
    module.evaluate_artifact(case, template, inspect_presentation(str(template)), checks)
    assert not next(x for x in checks if x["id"] == "required_layout_names")["passed"]
    assert not next(x for x in checks if x["id"] == "required_text")["passed"]


def test_page_compiler_clears_group_samples_and_notes(template, tmp_path):
    spec = spec_for(template)
    spec["pages"][0]["text"][0].update(font_size_pt=18, alignment="left", auto_size="none", word_wrap=True, font_family="WenQuanYi Zen Hei", character_spacing_pt=0)
    spec_path = tmp_path / "presentation-plan.json"
    spec_path.write_text(json.dumps(spec))
    result = build(spec_path, tmp_path / "run", False)
    inv = inspect_presentation(result["artifact"])
    assert len(inv["slides"]) == 1
    assert all(x["passed"] for x in content_checks(inv, spec["requirements"], spec["forbidden_text"]))
    assert not inv["slides"][0]["notes"]
    slide = Presentation(result["artifact"]).slides[0]
    from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
    assert slide.shapes[0].text_frame.auto_size == MSO_AUTO_SIZE.NONE
    assert slide.shapes[0].text_frame.paragraphs[0].alignment == PP_ALIGN.LEFT
    assert slide.shapes[0].text_frame.paragraphs[0].runs[0].font.size.pt == 18
    assert json.loads((tmp_path / "run/qa-report.json").read_text())["passed"] is False


def test_plan_rejects_changed_template_and_unknown_sources(template, tmp_path):
    spec = spec_for(template)
    spec["template"]["sha256"] = "0" * 64
    with pytest.raises(ValueError, match="version"):
        compile_presentation(spec, tmp_path)
    spec = spec_for(template)
    spec["pages"][0]["text"][0]["source_ids"] = ["missing"]
    with pytest.raises(ValueError, match="source id"):
        compile_presentation(spec, tmp_path)


def test_qa_requires_complete_review_and_current_hashes(template, tmp_path):
    image = tmp_path / "page.png"
    from PIL import Image
    Image.new("RGB", (40, 30)).save(image)
    manifest = {"rendered": True, "source": {"sha256": sha256_file(template), "slide_count": 1}, "slides": [{"slide_index": 0, "output": str(image), "sha256": sha256_file(image)}]}
    review = {"schema": "powerpoint.visual-review/v1", "artifact_sha256": sha256_file(template), "reviewer": "test fixture", "slides": []}
    inv = inspect_presentation(str(template))
    assert not quality_report(inv, template, manifest=manifest, review=review)["passed"]
    review["slides"] = [{"slide_index": 0, "image_sha256": sha256_file(image), "observations": "Synthetic QA validator fixture", "checks": {k: "pass" for k in VISUAL_CHECKS}}]
    assert quality_report(inv, template, manifest=manifest, review=review)["passed"]
    review["slides"][0]["checks"]["text_fit"] = "fail"
    assert not quality_report(inv, template, manifest=manifest, review=review)["passed"]
    image.write_bytes(b"changed")
    assert not quality_report(inv, template, manifest=manifest, review=review)["render_verified"]


def test_material_cell_sources_and_missing_formula_cache(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"] = "Revenue (万元)"
    sheet["B1"] = 120
    sheet["B2"] = "=B1*2"
    path = tmp_path / "facts.xlsx"
    workbook.save(path)
    manifest = extract_materials([path], tmp_path / "assets")
    assert not manifest["ready"]
    assert any(x["code"] == "formula_cache_missing" for x in manifest["issues"])
    missing = next(x for x in manifest["records"] if x["location"]["cell"] == "B2")
    assert missing["value"] is None and missing["formula"] == "=B1*2"
    assert manifest["files"][0]["sha256"] == sha256_file(path)


def test_material_docx_order_and_blank_pdf(tmp_path):
    docx = pytest.importorskip("docx")
    pypdf = pytest.importorskip("pypdf")
    doc = docx.Document()
    doc.add_paragraph("First fact")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Second fact"
    doc.add_paragraph("Third fact")
    path = tmp_path / "facts.docx"
    doc.save(path)
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    pdf = tmp_path / "scan.pdf"
    with pdf.open("wb") as handle:
        writer.write(handle)
    result = extract_materials([path, pdf], tmp_path / "assets")
    assert [r["text"] for r in result["records"]] == ["First fact", "Second fact", "Third fact"]
    assert not result["ready"] and any(i["code"] == "no_extractable_pdf_text" for i in result["issues"])


def test_template_selection_filters_photo_and_chart_constraints():
    photo = select_templates({"pages": [{"id": "photo", "photos": 2}]})
    assert [x["id"] for x in photo["candidates"]] == ["showeet-startup"]
    chart = select_templates({"pages": [{"id": "chart", "chart": "COLUMN_CLUSTERED"}]})
    assert [x["id"] for x in chart["candidates"]] == ["deckary-qbr"]
    impossible = select_templates({"pages": [{"id": "chart1", "chart": "COLUMN_CLUSTERED"}, {"id": "chart2", "chart": "COLUMN_CLUSTERED"}]})
    assert impossible["matched"] is False


def test_compile_rejects_a_changed_material_file(template, tmp_path):
    source = tmp_path / "facts.md"
    source.write_text("Verified source")
    material = extract_materials([source], tmp_path / "assets")
    path = tmp_path / "materials.json"
    path.write_text(json.dumps(material))
    spec = spec_for(template)
    spec["materials"] = {"path": str(path), "sha256": sha256_file(path)}
    source.write_text("Changed fact")
    with pytest.raises(ValueError, match="source material changed"):
        compile_presentation(spec, tmp_path)


def test_failed_render_cannot_leave_a_passed_run(template, tmp_path, monkeypatch):
    import pptx_workflow
    path = tmp_path / "presentation-plan.json"
    path.write_text(json.dumps(spec_for(template)))
    run_dir = tmp_path / "run"
    build(path, run_dir, False)
    def broken(*args):
        raise RuntimeError("renderer unavailable")
    monkeypatch.setattr(pptx_workflow, "render", broken)
    with pytest.raises(RuntimeError, match="renderer unavailable"):
        build(path, run_dir, True)
    state = json.loads((run_dir / "run.json").read_text())
    assert state["status"] == "render_failed" and state["visual_review_pending"]
    assert list((run_dir / "revisions").glob("*/output.pptx"))


@pytest.mark.parametrize("fit", ["cover", "contain"])
def test_empty_picture_placeholder_insertion_and_aspect_fit(tmp_path, fit):
    from PIL import Image
    from pptx_from_complete_template import replace_picture
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[8])
    placeholder = slide.placeholders[1]
    bounds = (placeholder.left, placeholder.top, placeholder.width, placeholder.height)
    path = tmp_path / "wide.png"
    Image.new("RGB", (400, 100), "blue").save(path)
    replace_picture(placeholder, str(path), fit)
    output = tmp_path / "picture.pptx"
    prs.save(output)
    picture = Presentation(output).slides[0].placeholders[1]
    assert picture.image.size == (400, 100)
    if fit == "cover":
        assert tuple(getattr(picture, k) for k in ("left", "top", "width", "height")) == bounds
        assert picture.crop_left > 0 and picture.crop_right > 0
    else:
        assert abs(picture.width / picture.height - 4) < .01
        assert picture.crop_left == picture.crop_right == 0


def test_chart_units_and_duplicate_constraints_roundtrip(tmp_path):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx_from_complete_template import fill_chart, execute_plan
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.oxml.ns import qn
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["Q1", "Q2"]
    data.add_series("Old USD", [0, 0])
    frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(7), Inches(4), data)
    unit = OxmlElement("c:dispUnits")
    builtin = OxmlElement("c:builtInUnit")
    builtin.set("val", "millions")
    unit.append(builtin)
    frame.chart.value_axis._element.append(unit)
    fill_chart(frame, {"categories": ["Q1", "Q2"], "series": {"营收（万元）": [120, 150]}, "number_format": '0"万元"', "display_units": "none"})
    path = tmp_path / "chart.pptx"
    prs.save(path)
    chart = Presentation(path).slides[0].shapes[0].chart
    assert list(chart.series[0].values) == [120, 150]
    assert chart.value_axis.tick_labels.number_format == '0"万元"'
    assert chart.value_axis._element.find(qn("c:dispUnits")) is None
    spec = spec_for(path)
    spec["pages"].append(dict(spec["pages"][0], id="copy"))
    spec["pages"][0].pop("text")
    spec["pages"][1].pop("text")
    spec["pages"][0]["preserve"] = [inspect_presentation(str(path))["slides"][0]["shapes"][0]["recommended_ref"]]
    spec["pages"][1]["preserve"] = spec["pages"][0]["preserve"]
    with pytest.raises(ValueError, match="chart page cannot be reused"):
        compile_presentation(spec, tmp_path)
